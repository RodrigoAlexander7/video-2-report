#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Helpers de formato que escriben DENTRO de una celda de tabla Word.

Las plantillas institucionales (EPIS / UNSA) son una tabla grande cuyas filas
("Ejercicios resueltos", "Cuestionario", "Conclusiones", ...) se rellenan con
contenido. Este módulo concentra los helpers que antes estaban copiados en
cada generador de reportes. Hay dos familias, con estética distinta:

1. Familia "configurable" (tutoriales Unity): el tipo de letra, los tamaños, el
   interlineado y las sangrías salen de un objeto ``cfg`` (p. ej.
   ``config.ReportConfig``); los colores corporativos son constantes de este
   módulo.

       apply_cell_margins, add_cell_heading_1, add_cell_heading_2,
       add_cell_body_paragraph, add_cell_bullet_point, add_cell_code_block,
       add_cell_image_with_caption

2. Familia "compacta" (informe corto EPIS Negocios Electrónicos): Calibri con
   tamaños fijos, sin ``cfg``.

       add_cell_heading, add_cell_p, add_cell_bullet, add_cell_table,
       add_cell_placeholder

Los equivalentes que escriben sobre el documento completo (no sobre una celda)
siguen en ``helpers.py``.
"""

import os
from typing import Protocol

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from .helpers import set_cell_border, set_cell_margins, set_cell_shading

# ---------------------------------------------------------------------------
# Estilo sobrio de la familia configurable, igual que los informes Typst: encabezados y viñetas en el
# color del texto (sin azules ni acentos), bloques de código en gris claro con borde fino y pies de figura
# con la etiqueta en negrita. Nada de recuadros de color: ni rojo para advertencias ni pastel para notas.
# ---------------------------------------------------------------------------
COLOR_BODY_TEXT = RGBColor(0x1F, 0x24, 0x21)       # Grafito suave para lectura (también encabezados y viñetas)
COLOR_CODE_TEXT = RGBColor(0x24, 0x29, 0x2E)
COLOR_BORDER = "D6D6D6"                            # Borde fino de los bloques de código (luma 214)
COLOR_CODE_BG = "F7F7F7"                           # Fondo del bloque de código (luma 247)


class CellStyleConfig(Protocol):
    """Atributos que la familia configurable lee de ``cfg`` (p. ej. ReportConfig)."""

    font_family_body: str
    font_family_code: str
    font_size_body_pt: float
    font_size_heading1_pt: float
    font_size_heading2_pt: float
    font_size_code_pt: float
    font_size_caption_pt: float
    body_line_spacing: float
    paragraph_space_after_pt: float
    bullet_indent_inches: float
    bullet_sub_indent_inches: float


# ===========================================================================
# Familia configurable (cfg): tutoriales Unity
# ===========================================================================

def apply_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Márgenes internos de respiración de una celda (en dxa).

    Igual que ``set_cell_margins`` pero con los márgenes por defecto amplios
    que usan las filas de contenido de las plantillas.
    """
    set_cell_margins(cell, top=top, bottom=bottom, left=left, right=right)


def add_cell_heading_1(cell, text, cfg: CellStyleConfig):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = cfg.font_family_body
    run.font.size = Pt(cfg.font_size_heading1_pt)
    run.font.color.rgb = COLOR_BODY_TEXT
    return p


def add_cell_heading_2(cell, text, cfg: CellStyleConfig):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = cfg.font_family_body
    run.font.size = Pt(cfg.font_size_heading2_pt)
    run.font.color.rgb = COLOR_BODY_TEXT
    return p


def add_cell_body_paragraph(cell, text, cfg: CellStyleConfig, bold_prefix="", italic=False):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(cfg.paragraph_space_after_pt)
    p.paragraph_format.line_spacing = cfg.body_line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = cfg.font_family_body
        r_pre.font.size = Pt(cfg.font_size_body_pt)
        r_pre.font.color.rgb = COLOR_BODY_TEXT

    run = p.add_run(text)
    run.font.name = cfg.font_family_body
    run.font.size = Pt(cfg.font_size_body_pt)
    run.font.color.rgb = COLOR_BODY_TEXT
    run.italic = italic
    return p


def add_cell_bullet_point(cell, bold_text, desc_text, cfg: CellStyleConfig, level: int = 1):
    p = cell.add_paragraph()
    indent = cfg.bullet_indent_inches if level == 1 else cfg.bullet_sub_indent_inches
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = cfg.body_line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    bullet_symbol = "•  " if level == 1 else "–  "
    r_sym = p.add_run(bullet_symbol)
    r_sym.font.name = cfg.font_family_body
    r_sym.font.size = Pt(cfg.font_size_body_pt)
    r_sym.font.color.rgb = COLOR_BODY_TEXT

    if bold_text:
        r_b = p.add_run(bold_text)
        r_b.bold = True
        r_b.font.name = cfg.font_family_body
        r_b.font.size = Pt(cfg.font_size_body_pt)
        r_b.font.color.rgb = COLOR_BODY_TEXT

    r_d = p.add_run(desc_text)
    r_d.font.name = cfg.font_family_body
    r_d.font.size = Pt(cfg.font_size_body_pt)
    r_d.font.color.rgb = COLOR_BODY_TEXT
    return p


def add_cell_code_block(cell, code_str, cfg: CellStyleConfig):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)

    pPr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{COLOR_CODE_BG}"/>')
    pPr.append(shd)

    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="4" w:space="8" w:color="{COLOR_BORDER}"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="4" w:color="{COLOR_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="4" w:color="{COLOR_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="4" w:color="{COLOR_BORDER}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

    run = p.add_run(code_str)
    run.font.name = cfg.font_family_code
    run.font.size = Pt(cfg.font_size_code_pt)
    run.font.color.rgb = COLOR_CODE_TEXT
    return p


def add_cell_image_with_caption(cell, img_path, caption_text, cfg: CellStyleConfig, width=Inches(5.4)):
    p_img = cell.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.add_run().add_picture(img_path, width=width)

    p_cap = cell.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    # "Figura 5:" en negrita y el resto normal, como los pies de figura de Typst (sin cursiva ni gris)
    label, sep, rest = caption_text.partition(":")
    parts = [(label + sep, True), (rest, False)] if sep and len(label) <= 14 else [(caption_text, False)]
    for text, bold in parts:
        if not text:
            continue
        r_cap = p_cap.add_run(text)
        r_cap.font.name = cfg.font_family_body
        r_cap.font.size = Pt(cfg.font_size_caption_pt)
        r_cap.bold = bold
        r_cap.font.color.rgb = COLOR_BODY_TEXT


def add_cell_code_figure(cell, snippet_id, evidence_dir, caption_text, cfg: CellStyleConfig,
                         text_pt=9.5, max_width_in=6.0):
    """Inserta la imagen de un fragmento de código (``main.py code-image``) por su id del manifiesto.

    El ancho se calcula para que el texto del código mida ``text_pt`` en el documento sea cual sea el fragmento
    (acotado a ``max_width_in``); así todos los fragmentos de un informe se ven con el mismo tamaño de letra.
    """
    from ..editor_engine.render import FONT_PX
    from ..evidence_engine import manifest as mf
    item = mf.load_manifest(str(evidence_dir)).get("items", {}).get(snippet_id)
    if not item or item.get("source") != "code":
        raise KeyError(f"'{snippet_id}' no es un fragmento de código registrado en {evidence_dir}/manifest.json "
                       "(genera las imágenes con: python main.py code-image <caso>)")
    path = os.path.join(str(evidence_dir), item["image"])
    if mf.sha256_file(path) != item["image_sha256"]:
        raise ValueError(f"la imagen de '{snippet_id}' fue modificada después de generarse: regénerala")
    natural_in = item["width_px"] / 96                       # ancho al 100 % (texto de FONT_PX px = FONT_PX*0.75 pt)
    width_in = min(max_width_in, natural_in * text_pt / (FONT_PX * 0.75))
    add_cell_image_with_caption(cell, path, caption_text, cfg, width=Inches(width_in))


# ===========================================================================
# Familia compacta (Calibri, tamaños fijos): informe corto EPIS Negocios Electrónicos
# ===========================================================================

def add_cell_heading(cell, text, level=1):
    p = cell.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return p


def add_cell_p(cell, text, bold_prefix="", italic_prefix=False, space_after=4):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        if italic_prefix:
            r_pre.italic = True
    if text:
        r_txt = p.add_run(text)
        r_txt.font.name = 'Calibri'
        r_txt.font.size = Pt(10)
    return p


def add_cell_bullet(cell, text, bold_prefix=""):
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    r_bullet = p.add_run("▪  ")
    r_bullet.bold = True
    r_bullet.font.name = 'Calibri'
    r_bullet.font.size = Pt(10)
    r_bullet.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
    if text:
        r_txt = p.add_run(text)
        r_txt.font.name = 'Calibri'
        r_txt.font.size = Pt(10)
    return p


def add_cell_table(cell, headers, data, col_widths=None):
    tbl = cell.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Cabecera
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "333333")
        set_cell_border(hdr_cells[i], color="333333", sz="6")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Filas de datos
    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_border(row_cells[c_idx], color="D3D3D3", sz="4")
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=80, right=80)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(8.5)

    if col_widths:
        for row in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)

    p_spacer = cell.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(4)


def add_cell_placeholder(cell, fig_num, title, module, objective, test_data, visible_items, expected_result):
    tbl = cell.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    c.width = Inches(6.4)

    set_cell_shading(c, "F8F9FA")
    set_cell_border(c, left='single', right='single', top='single', bottom='single', color='555555', sz='8')
    set_cell_margins(c, top=120, bottom=120, left=150, right=150)

    p0 = c.paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(3)
    r_badge = p0.add_run(f"EVIDENCIA DE PROTOTIPO - FIGURA {fig_num}: {title.upper()}\n")
    r_badge.bold = True
    r_badge.font.name = 'Calibri'
    r_badge.font.size = Pt(10)
    r_badge.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    p1 = c.add_paragraph()
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15

    r1 = p1.add_run("• Módulo Odoo: ")
    r1.bold = True
    p1.add_run(f"{module}\n")

    r2 = p1.add_run("• Objetivo de la prueba: ")
    r2.bold = True
    p1.add_run(f"{objective}\n")

    r3 = p1.add_run("• Datos de prueba simulados: ")
    r3.bold = True
    p1.add_run(f"{test_data}\n")

    r4 = p1.add_run("• Elementos visuales requeridos en la captura: ")
    r4.bold = True
    p1.add_run(f"{visible_items}\n")

    r5 = p1.add_run("• Resultado esperado: ")
    r5.bold = True
    p1.add_run(f"{expected_result}")

    for r in p1.runs:
        r.font.name = 'Calibri'
        r.font.size = Pt(9)

    p2 = c.add_paragraph()
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after = Pt(4)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_box = p2.add_run(
        "┌────────────────────────────────────────────────────────────────────────┐\n"
        "│                                                                        │\n"
        "│     [ PEGAR AQUÍ LA CAPTURA DE PANTALLA DE ODOO COMMUNITY UNA VEZ      │\n"
        "│                    EJECUTADA ESTA CONFIGURACIÓN ]                      │\n"
        "│                                                                        │\n"
        "└────────────────────────────────────────────────────────────────────────┘"
    )
    r_box.font.name = 'Consolas'
    r_box.font.size = Pt(8)
    r_box.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p_cap = cell.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(6)
    rcap = p_cap.add_run(f"Figura {fig_num}. {title}. Fuente: Elaboración propia en Odoo Community Edition.")
    rcap.font.size = Pt(8.5)
    rcap.font.name = 'Calibri'
    rcap.italic = True
    rcap.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
