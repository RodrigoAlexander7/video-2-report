# -*- coding: utf-8 -*-
"""Genera el .docx del reporte de Unity sobre la plantilla institucional (tabla EPIS) a partir de ``report.toml``."""

from __future__ import annotations

import copy
import re
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph

from config import ReportConfig
from ..docx_engine.block_rows import BlockRows
from ..docx_engine.cell_helpers import (add_cell_heading_1, add_cell_heading_2, add_cell_image_with_caption,
                                        add_cell_runs_paragraph, apply_cell_margins)
from ..evidence_engine import manifest as mf
from . import code as codemod
from .checks import check_lab, format_scorecard
from .figures import FigureError, FigureResult, capture_figures
from .layout import fit_capture, fit_code, prepare_bitmap
from .model import Block, Lab, ReportError, load_lab

REPO_ROOT = Path(__file__).resolve().parents[2]
_MARKUP = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\{fig:[\w\-]+\})")


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", text.upper())
    return re.sub(r"\s+", " ", "".join(c for c in text if not unicodedata.combining(c))).strip()


def _unique_cells(row) -> List[Tuple[int, Any]]:
    seen, out = [], []
    for i, cell in enumerate(row.cells):
        if cell._tc in seen:
            continue
        seen.append(cell._tc)
        out.append((i, cell))
    return out


def locate(table) -> Dict[str, int]:
    """Filas de la plantilla por su rótulo (no por posición): distintas plantillas EPIS mueven las filas."""
    found: Dict[str, int] = {}
    wanted = {"course": "ASIGNATURA", "title": "TITULO DE LA PRACTICA", "number": "NUMERO DE PRACTICA", "date": "FECHA DE PRESENTACION",
              "members": "INTEGRANTE", "instructor": "DOCENTE", "exercises": "EJERCICIOS RESUELTOS", "questions": "CUESTIONARIO",
              "conclusions_h": "CONCLUSIONES", "references_h": "REFERENCIAS"}
    for i, row in enumerate(table.rows):
        label = _norm(_unique_cells(row)[0][1].text)
        for key, prefix in wanted.items():
            if key not in found and label.startswith(prefix):
                found[key] = i
    missing = [k for k in ("course", "title", "number", "date", "members", "instructor", "exercises", "questions", "conclusions_h", "references_h")
               if k not in found]
    if missing:
        raise ReportError([f"la plantilla no tiene las filas esperadas ({', '.join(missing)}): "
                           "usa la plantilla EPIS de la guía (tabla con INFORMACIÓN BÁSICA, EJERCICIOS RESUELTOS, CUESTIONARIO, "
                           "CONCLUSIONES y REFERENCIAS)"])
    found["conclusions"] = found["conclusions_h"] + 1
    found["references"] = found["references_h"] + 1
    if found["references"] >= len(table.rows) or found["conclusions"] >= len(table.rows):
        raise ReportError(["la plantilla no tiene fila de contenido bajo CONCLUSIONES o REFERENCIAS"])
    return found


def _set_text(cell, text: str) -> None:
    """Escribe ``text`` conservando el formato del primer tramo de la celda."""
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]
    runs = [r for r in p.runs if r.text.strip()] or p.runs
    if runs:
        keep = runs[0]._r
        keep_run = runs[0]
        keep_run.text = text
        for r in p.runs:
            if r._r is not keep:
                r._r.getparent().remove(r._r)
    else:
        p.add_run(text)


def _value_after(row, label_prefix: str):
    cells = _unique_cells(row)
    for k, (_, cell) in enumerate(cells):
        if _norm(cell.text).startswith(label_prefix) and k + 1 < len(cells):
            return cells[k + 1][1]
    return None


def _fill_info(table, rows: Dict[str, int], meta: Dict[str, Any]) -> None:
    row = lambda key: table.rows[rows[key]]
    cells = lambda key: _unique_cells(row(key))
    if meta.get("course"):
        _set_text(cells("course")[1][1], meta["course"])
    _set_text(cells("title")[1][1], meta["title"])
    _set_text(cells("number")[1][1], str(meta["number"]))
    if meta.get("term") and (c := _value_after(row("number"), "ANO LECTIVO")):
        _set_text(c, meta["term"])
    if meta.get("level") and (c := _value_after(row("number"), "NRO. SEMESTRE")):
        _set_text(c, meta["level"])
    _set_text(cells("date")[1][1], meta["date"])
    if meta.get("time") and (c := _value_after(row("date"), "HORA DE PRESENTACION")):
        _set_text(c, meta["time"])
    # Integrantes: la celda tiene el rótulo y debajo los nombres
    cell = cells("members")[0][1]
    names = meta["members"]
    paras = cell.paragraphs
    for extra in paras[2:]:
        extra._p.getparent().remove(extra._p)
    if len(paras) < 2:
        cell.add_paragraph()
    _set_text_in_paragraph(cell.paragraphs[1], names[0])
    last = cell.paragraphs[1]._p
    for name in names[1:]:
        clone = copy.deepcopy(last)
        last.addnext(clone)
        last = clone
        _set_text_in_paragraph(Paragraph(clone, cell), name)
    if meta.get("instructor"):
        cell = cells("instructor")[0][1]
        if len(cell.paragraphs) < 2:
            cell.add_paragraph()
        _set_text_in_paragraph(cell.paragraphs[1], meta["instructor"])


def _set_text_in_paragraph(p, text: str) -> None:
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)


def _clear(cell) -> None:
    for t in list(cell.tables):
        t._tbl.getparent().remove(t._tbl)
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in list(cell.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    apply_cell_margins(cell)


def _finish(cell) -> None:
    """Quita el párrafo vacío inicial que deja `_clear` cuando ya hay contenido."""
    first = cell.paragraphs[0]
    if len(cell.paragraphs) > 1 and not first.text.strip() and not first._p.xpath(".//w:drawing"):
        first._p.getparent().remove(first._p)


def segments(text: str, fig_no: Dict[str, int]) -> List[Tuple[str, str]]:
    """Texto con ``**negrita**``, `` `código` `` y ``{fig:id}`` -> tramos (texto, estilo)."""
    out: List[Tuple[str, str]] = []
    for part in _MARKUP.split(text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], "bold"))
        elif part.startswith("`"):
            out.append((part[1:-1], "code"))
        elif part.startswith("{fig:"):
            out.append((f"Figura {fig_no[part[5:-1]]}", ""))
        else:
            out.append((part, ""))
    return out


class _Ctx:
    def __init__(self, lab: Lab, cfg: ReportConfig, figures_dir: Path, evidence_dir: Path, results: Dict[str, FigureResult],
                 code_plans: Dict[str, List[Any]]):
        self.lab, self.cfg, self.figures_dir, self.evidence_dir = lab, cfg, figures_dir, evidence_dir
        self.table = None                          # tabla institucional; se asigna al abrir la plantilla
        self.results, self.code_plans = results, code_plans
        self.fig_no: Dict[str, int] = {}
        n = 0
        for b in lab.blocks:
            if b.figure:
                n += 1
                self.fig_no[b.figure.id] = n
            if b.code:
                self.fig_no[b.code.id] = n + 1
                n += len(code_plans[b.code.id])
        self.docx_dir = figures_dir / "_docx"
        self.next_no = 0

    def take_no(self) -> int:
        self.next_no += 1
        return self.next_no


def _add_figure(cell, ctx: _Ctx, image: Path, caption: str, fit, photographic: bool = False) -> None:
    prepared = prepare_bitmap(image, ctx.docx_dir / image.stem, fit.width_in, photographic=photographic)
    add_cell_image_with_caption(cell, str(prepared), f"Figura {ctx.take_no()}: {caption}", ctx.cfg, width=Inches(fit.width_in))


def _add_video_figure(cell, ctx: _Ctx, block: Block) -> None:
    fs = block.figure
    res = ctx.results[fs.id]
    fit = fit_capture(res.width, res.height, res.ui_text_px, res.legible, f"figura {fs.id}")
    _add_figure(cell, ctx, ctx.figures_dir / res.path, fs.caption, fit, photographic=not fs.legible and fs.kind != "gif")


def _add_code(cell, ctx: _Ctx, block: Block) -> None:
    spec = block.code
    snips = ctx.code_plans[spec.id]
    for k, snip in enumerate(snips, 1):
        item = codemod.ensure_image(snip, ctx.evidence_dir)
        fit = fit_code(item["width_px"], item["height_px"], f"código {snip.id}")
        a, b = item["ranges"][0][0], item["ranges"][-1][1]
        if spec.caption:
            caption = spec.caption if k == 1 else f"{spec.caption.rstrip('.')} (continuación)."
        else:
            caption = f"Fragmento de {Path(spec.file).name}, líneas {a} a {b}."
        _add_figure(cell, ctx, ctx.evidence_dir / item["image"], caption, fit)


def _compose_exercises(row, ctx: _Ctx) -> None:
    """Cada grupo (encabezados + paso + figura/código) va en su propia fila indivisible: si no cabe en la página
    que queda, pasa entero a la siguiente (sin encabezados huérfanos ni figuras separadas de su texto)."""
    lab, cfg = ctx.lab, ctx.cfg
    rows = BlockRows(ctx.table, row, _clear)
    cell = rows.next_cell()
    add_cell_heading_1(cell, "EJERCICIOS RESUELTOS:", cfg)
    if lab.meta.get("video_link"):
        add_cell_runs_paragraph(cell, [("Video de la práctica: ", "bold"), (lab.meta["video_link"], "")], cfg, justify=False)
    h1 = h2 = 0
    pending: List[Block] = []                       # encabezados que esperan a su primer paso
    for b in lab.blocks:
        if b.kind == "heading":
            pending.append(b)
            continue
        group, pending = pending + [b], []
        _finish(cell)
        cell = rows.next_cell()
        for g in group:
            if g.kind == "heading":
                if g.level == 1:
                    h1, h2 = h1 + 1, 0
                    add_cell_heading_2(cell, f"{h1}. {g.text}", cfg)
                else:
                    h2 += 1
                    add_cell_runs_paragraph(cell, [(f"{h1}.{h2} {g.text}", "bold")], cfg, justify=False)
                continue
            if g.kind == "step":
                add_cell_runs_paragraph(cell, segments(g.text, ctx.fig_no), cfg)
            if g.figure:
                _add_video_figure(cell, ctx, g)
            if g.code:
                _add_code(cell, ctx, g)
        _finish(cell)
    rows.finish()


def _compose_questions(row, ctx: _Ctx) -> None:
    cfg = ctx.cfg
    rows = BlockRows(ctx.table, row, _clear)
    cell = rows.next_cell()
    add_cell_heading_1(cell, "CUESTIONARIO:", cfg)
    for i, q in enumerate(ctx.lab.questions, 1):
        if i > 1:
            _finish(cell)
            cell = rows.next_cell()
        add_cell_runs_paragraph(cell, [(f"{i}. ", "bold")] + [(t, "bold" if s != "code" else "code") for t, s in segments(q["q"], ctx.fig_no)],
                                cfg, justify=False)
        for para in re.split(r"\n\s*\n", q["a"]):
            add_cell_runs_paragraph(cell, segments(para.strip(), ctx.fig_no), cfg)
    _finish(cell)
    rows.finish()


def _compose_conclusions(cell, ctx: _Ctx) -> None:
    _clear(cell)
    for text in ctx.lab.conclusions:
        p = add_cell_runs_paragraph(cell, [("•  ", "")] + segments(text, ctx.fig_no), ctx.cfg)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.2)
    _finish(cell)


def _compose_references(cell, ctx: _Ctx) -> None:
    _clear(cell)
    for n, text in enumerate(ctx.lab.references, 1):
        p = add_cell_runs_paragraph(cell, [(f"[{n}]  ", "")] + segments(text, ctx.fig_no), ctx.cfg, justify=False, space_after=3)
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
    _finish(cell)


def default_dirs(project_dir: Path) -> Tuple[Path, Path, Path]:
    from ..core.registry import resolve_slug
    out = REPO_ROOT / "output" / resolve_slug(project_dir)
    return out, out / "figures", out / "evidence"


def build(project_dir, out_dir: Optional[str] = None, force: bool = False, strict: bool = False,
          only: Optional[List[str]] = None, cfg: Optional[ReportConfig] = None) -> Path:
    """Valida, captura las figuras, genera las imágenes de código y compone el .docx. Devuelve su ruta."""
    project_dir = Path(project_dir).resolve()
    lab = load_lab(project_dir)
    out, figures_dir, evidence_dir = default_dirs(project_dir)
    if out_dir:
        out = Path(out_dir)
        figures_dir, evidence_dir = out / "figures", out / "evidence"
    cfg = cfg or ReportConfig()
    video = project_dir / lab.meta["video"] if lab.meta.get("video") else None

    # 1) revisión previa (texto, código, citas): falla antes de gastar tiempo en el video
    pre = check_lab(lab)
    if pre.errors:
        raise ReportError([str(i) for i in pre.errors])
    # 2) figuras del video (con caché)
    results: Dict[str, FigureResult] = {}
    if lab.figures:
        results = capture_figures(lab, video, figures_dir, evidence_dir, force=force, only=only, langs=lab.meta.get("ocr_langs", "eng"))
    # 3) revisión con tamaños reales de las figuras
    rep = check_lab(lab, results)
    if rep.errors or (strict and rep.warnings):
        raise ReportError([str(i) for i in (rep.errors + (rep.warnings if strict else []))])
    for issue in rep.warnings:
        print(issue)
    print(format_scorecard(rep))

    code_plans = {b.code.id: codemod.plan_code(b.code, project_dir, lab.meta.get("theme", "dark")) for b in lab.blocks if b.code}
    ctx = _Ctx(lab, cfg, figures_dir, evidence_dir, results, code_plans)

    document = docx.Document(str(project_dir / lab.meta["template"]))
    if not document.tables:
        raise ReportError(["la plantilla no contiene la tabla institucional"])
    table = document.tables[0]
    rows = locate(table)
    _fill_info(table, rows, lab.meta)
    # las filas se toman antes de componer: clonar filas desplaza los índices de las siguientes
    r_ex, r_q, r_c, r_r = (table.rows[rows[k]] for k in ("exercises", "questions", "conclusions", "references"))
    ctx.table = table
    _compose_exercises(r_ex, ctx)
    _compose_questions(r_q, ctx)
    _compose_conclusions(r_c.cells[0], ctx)
    _compose_references(r_r.cells[0], ctx)
    document.core_properties.title = f"Práctica {lab.meta['number']} — {lab.meta['title']}"
    document.core_properties.author = ", ".join(lab.meta["members"])
    path = out / f"Reporte_Practica_{lab.meta['number']}_Unity.docx"
    out.mkdir(parents=True, exist_ok=True)
    document.save(path)
    print(f"✓ {path}")
    return path
