import os
import sys
import argparse
import docx
from docx.shared import Pt

if not __package__:
    # Ejecución directa (python src/reports/.../archivo.py): exponer la raíz del proyecto
    # para resolver `config` y `src.*`. Con main.py o `python -m` ya está en sys.path.
    sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../../..")))

from config import ReportConfig
from src.docx_engine.cell_helpers import (
    apply_cell_margins,
    add_cell_heading_1,
    add_cell_heading_2,
    add_cell_body_paragraph,
    add_cell_bullet_point,
    add_cell_code_block,
    add_cell_image_with_caption,
)

def build_lab06_report(
    template_path: str,
    output_path: str,
    media_dir: str,
    cfg: ReportConfig = ReportConfig()
):
    print(f"Cargando plantilla desde: {template_path}")
    doc = docx.Document(template_path)
    table = doc.tables[0]

    # --- 1. ENCABEZADO Y METADATOS DE LA PRÁCTICA ---
    # Título de la práctica
    cell_tit = table.rows[2].cells[1]
    cell_tit.text = "Trabajo con Rotación y Time.DeltaTime"
    p_tit = cell_tit.paragraphs[0]
    p_tit.runs[0].font.name = cfg.font_family_body
    p_tit.runs[0].font.size = Pt(cfg.font_size_body_pt)
    p_tit.runs[0].bold = True

    # Número de práctica
    cell_num = table.rows[3].cells[1]
    cell_num.text = "06"
    p_num = cell_num.paragraphs[0]
    p_num.runs[0].font.name = cfg.font_family_body
    p_num.runs[0].font.size = Pt(cfg.font_size_body_pt)
    p_num.runs[0].bold = True

    # Fecha de presentación
    cell_fec = table.rows[4].cells[1]
    cell_fec.text = "12/10/2026"
    p_fec = cell_fec.paragraphs[0]
    p_fec.runs[0].font.name = cfg.font_family_body
    p_fec.runs[0].font.size = Pt(cfg.font_size_body_pt)

    # --- 2. SECCIÓN: RESULTADOS Y PRUEBAS (TUTORIAL PROGRESIVO POR ETAPAS) ---
    cell_ej = table.rows[9].cells[0]
    cell_ej.text = ""
    apply_cell_margins(cell_ej)

    add_cell_heading_1(cell_ej, "EJERCICIOS RESUELTOS: TUTORIAL DE CONTROLADOR DE PERSONAJE (CatControl)", cfg)

    add_cell_body_paragraph(cell_ej, 
        "El presente tutorial describe detalladamente el proceso experimental y de desarrollo llevado a cabo en Unity 6 "
        "para implementar la rotación orientada de un personaje tridimensional ('Mango'), resolver de manera rigurosa "
        "la dependencia de cuadros por segundo mediante Time.deltaTime, corregir errores de traslación física mediante "
        "AddRelativeForce y solucionar desalineaciones espaciales en los ejes locales del modelado 3D. "
        "La guía está organizada en 5 etapas progresivas y secuenciales, incorporando código incremental y evidencia visual "
        "(capturas de alta definición y animaciones GIF) en cada paso y subpaso.",
        cfg
    )

    # =========================================================================
    # ETAPA 1
    # =========================================================================
    add_cell_heading_2(cell_ej, "Etapa 1: Análisis del Estado Previo e Implementación de la Rotación Angular sobre el Eje Y", cfg)
    add_cell_body_paragraph(cell_ej, 
        "En esta etapa se evalúa la mecánica de desplazamiento previa y se programa el control de giro angular del personaje "
        "en función de las teclas laterales del usuario.",
        cfg
    )

    # Subpaso 1.1
    add_cell_bullet_point(cell_ej, "Paso 1.1: Diagnóstico del Movimiento Cartesiano Previo. ", 
        "En la iteración previa del laboratorio, el personaje se desplazaba sumando directamente desplazamientos discretos sobre las coordenadas "
        "cartesianas globales (X, Y, Z). Si bien permitía el avance y el salto, el modelo del gato permanecía inmóvil en su orientación angular, "
        "desplazándose de manera rígida y poco natural sin girar hacia la dirección de avance.", 
        cfg, level=1)
    
    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig1_prev_state.png"),
        "Figura 1: Estado inicial del personaje en Unity evidenciando traslación cartesiana estática sin orientación angular.",
        cfg)

    # Subpaso 1.2
    add_cell_bullet_point(cell_ej, "Paso 1.2: Declaración de Variables de Control y Referencia a Rigidbody. ", 
        "Para migrar a un control dinámico y angular, se abre el script C# 'CatControl.cs' y se declaran las variables públicas 'thrust' (fuerza de impulso = 5.0f) "
        "y 'rotationSpeed' (velocidad angular = 50.0f). En el ciclo 'Start()', se captura la instancia física del componente Rigidbody asignado al GameObject "
        "mediante la instrucción 'GetComponent<Rigidbody>()':", 
        cfg, level=1)

    code_step1_2 = """// CatControl.cs - Fragmento Incremental: Variables y Captura de Componentes
using UnityEngine;

public class CatControl : MonoBehaviour
{
    public float thrust = 5.0f;
    public float rotationSpeed = 50.0f;

    Rigidbody rb;

    void Start()
    {
        rb = GetComponent<Rigidbody>();
    }

    void Update()
    {
        ProcesarInput();
    }
}"""
    add_cell_code_block(cell_ej, code_step1_2, cfg)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig2_code_variables.png"),
        "Figura 2: Declaración de variables de velocidad angular (rotationSpeed), empuje (thrust) y captura del Rigidbody en Start().",
        cfg)

    # Subpaso 1.3
    add_cell_bullet_point(cell_ej, "Paso 1.3: Implementación de Rotación en el Eje Vertical mediante transform.Rotate. ", 
        "En el método 'ProcesarInput()', se captura la pulsación de las teclas 'D' y 'A' para aplicar rotación. Dado que el personaje se encuentra apoyado "
        "sobre el plano del terreno, girar a la derecha o izquierda corresponde a rotar en torno a su eje vertical Y. En Unity, una rotación positiva en Y "
        "(Vector3.up) vista desde arriba gira en sentido horario (hacia la derecha, tecla 'D'), mientras que Vector3.down (o -Vector3.up) gira en sentido "
        "antihorario (hacia la izquierda, tecla 'A'). Nótese el uso directo de la propiedad 'transform' expuesta por MonoBehaviour sin invocar GetComponent:", 
        cfg, level=1)

    code_step1_3 = """// CatControl.cs - Fragmento Incremental: Método ProcesarInput() con Rotación
    private void ProcesarInput()
    {
        if (Input.GetKey(KeyCode.Space))
        {
            rb.AddRelativeForce(Vector3.up * thrust);
        }

        if (Input.GetKey(KeyCode.D))
        {
            transform.Rotate(Vector3.up * rotationSpeed);
        }
        else if (Input.GetKey(KeyCode.A))
        {
            transform.Rotate(Vector3.down * rotationSpeed);
        }
    }"""
    add_cell_code_block(cell_ej, code_step1_3, cfg)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig3_code_rotate.png"),
        "Figura 3: Invocación directa a transform.Rotate() sobre Vector3.up y Vector3.down según la entrada de teclado.",
        cfg)

    # =========================================================================
    # ETAPA 2
    # =========================================================================
    add_cell_heading_2(cell_ej, "Etapa 2: Diagnóstico Experimental de la Dependencia de Cuadros (Efecto 'Gato Poseído')", cfg)
    add_cell_body_paragraph(cell_ej, 
        "En esta etapa se comprueban los efectos inmediatos de ejecutar la rotación angular en Update() sin escalar por tiempo real.",
        cfg
    )

    # Subpaso 2.1
    add_cell_bullet_point(cell_ej, "Paso 2.1: Ejecución en Modo Play a Alta Tasa de Refresco. ", 
        "Al compilar el código e ingresar en modo Play en Unity, se presiona la tecla 'D' o 'A'. El resultado visual es caótico: el personaje comienza "
        "a girar a una velocidad exorbitante y desenfrenada ('efecto gato poseído'), perdiendo toda estabilidad y control motriz.", 
        cfg, level=1)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig4_crazy_spin.gif"),
        "Figura 4 (GIF Animado): Giro desenfrenado e incontrolable del personaje al carecer de independencia de frames.",
        cfg)

    # Subpaso 2.2
    add_cell_bullet_point(cell_ej, "Paso 2.2: Inspección de Métricas de Rendimiento en el Panel Stats. ", 
        "Para diagnosticar la causa raíz de este comportamiento, se abre el panel 'Statistics' de la ventana Game en Unity. Las métricas del motor revelan "
        "que la escena se ejecuta a una tasa de refresco de 229.9 FPS (Global Frametime = 4.4 ms). Dado que Update() se ejecuta una vez por cada cuadro "
        "renderizado, la llamada 'transform.Rotate(Vector3.up * 50)' aplica 50 grados en cada cuadro. A 230 FPS, esto produce una velocidad angular de "
        "50 * 230 = 11,500 grados por segundo (casi 32 vueltas completas de 360° en un solo segundo). En una computadora más lenta corriendo a 30 FPS, "
        "el gato giraría a tan solo 1,500 grados por segundo, evidenciando una dependencia crítica del hardware que destruye la experiencia de juego.", 
        cfg, level=1)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig5_stats_panel.png"),
        "Figura 5: Panel Statistics de Unity revelando una tasa de 229.9 FPS y frametime de 4.4 ms como origen del giro desmedido.",
        cfg)

    # =========================================================================
    # ETAPA 3
    # =========================================================================
    add_cell_heading_2(cell_ej, "Etapa 3: Estabilización Temporal e Independencia de Frames mediante Time.deltaTime", cfg)
    add_cell_body_paragraph(cell_ej, 
        "Esta etapa aborda la solución matemática al problema de desacoplamiento de cuadros por segundo utilizando la propiedad estática Time.deltaTime.",
        cfg
    )

    # Subpaso 3.1
    add_cell_bullet_point(cell_ej, "Paso 3.1: Incorporación del Factor Escalar Time.deltaTime. ", 
        "Para que la rotación se mida en grados por segundo en lugar de grados por cuadro, se multiplica la velocidad angular por 'Time.deltaTime'. "
        "Como 'Time.deltaTime' representa la fracción de segundo transcurrida desde el cuadro anterior (aproximadamente 0.00435 segundos a 230 FPS), "
        "la rotación acumulada en el lapso de un segundo será exactamente igual a 50 grados, sin importar si el motor renderiza a 30, 60 o 240 FPS:", 
        cfg, level=1)

    code_step3_1 = """// CatControl.cs - Fragmento Incremental: Multiplicación por Time.deltaTime
        if (Input.GetKey(KeyCode.D))
        {
            transform.Rotate(Vector3.up * rotationSpeed * Time.deltaTime);
        }
        else if (Input.GetKey(KeyCode.A))
        {
            transform.Rotate(Vector3.down * rotationSpeed * Time.deltaTime);
        }"""
    add_cell_code_block(cell_ej, code_step3_1, cfg)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig6_code_deltatime.png"),
        "Figura 6: Corrección del script aplicando el factor Time.deltaTime en las expresiones de rotación angular.",
        cfg)

    # Subpaso 3.2
    add_cell_bullet_point(cell_ej, "Paso 3.2: Validación de Rotación Suave y Homogénea en Modo Play. ", 
        "Tras guardar y recompilar el script, se vuelve a ejecutar la escena en Unity. Al mantener presionadas las teclas 'A' y 'D', el personaje gira "
        "de manera suave, uniforme y perfectamente controlable a una tasa constante de 50 grados por segundo, logrando la ansiada independencia de frames.", 
        cfg, level=1)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig7_smooth_rotation.gif"),
        "Figura 7 (GIF Animado): Rotación suave, estable y predecible del personaje tras la aplicación de Time.deltaTime.",
        cfg)

    # =========================================================================
    # ETAPA 4
    # =========================================================================
    add_cell_heading_2(cell_ej, "Etapa 4: Implementación del Desplazamiento Lineal (W/S) y Corrección de Ejes de Giro", cfg)
    add_cell_body_paragraph(cell_ej, 
        "En esta etapa se implementa el movimiento de avance y retroceso del personaje, analizando la diferencia fundamental entre rotación cinemática y traslación física.",
        cfg
    )

    # Subpaso 4.1
    add_cell_bullet_point(cell_ej, "Paso 4.1: Detección del Error de Cabeceo al Utilizar transform.Rotate en W y S. ", 
        "Inicialmente se intentó programar el avance ('W') y retroceso ('S') reutilizando 'transform.Rotate' con 'Vector3.left' y 'Vector3.right'. "
        "Al probar esta aproximación en Unity, el gato no avanzó sobre el suelo, sino que comenzó a realizar volteretas sobre su eje transversal (cabeceo / pitch). "
        "Esto ocurre porque 'Transform.Rotate' altera exclusivamente la orientación angular tridimensional del objeto y no traslada su centro de masa.", 
        cfg, level=1)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig8_pitch_spin_error.gif"),
        "Figura 8 (GIF Animado): Anomalía de cabeceo y vuelco sobre el eje transversal al usar transform.Rotate para traslación.",
        cfg)

    # Subpaso 4.2
    add_cell_bullet_point(cell_ej, "Paso 4.2: Sustitución por Fuerza Física Relativa (rb.AddRelativeForce). ", 
        "Para lograr que el personaje se desplace efectivamente en el espacio tridimensional hacia la dirección a la que apunta su frente, se descarta "
        "'transform.Rotate' en W/S y se utiliza el componente 'Rigidbody'. Se introduce la variable 'public float moveForce = 10f' (calibrada luego a 3f) "
        "y se invoca el método 'rb.AddRelativeForce()', el cual aplica una fuerza vectorial con respecto al sistema de coordenadas local del propio cuerpo rígido:", 
        cfg, level=1)

    code_step4_2 = """// CatControl.cs - Fragmento Incremental: Traslación mediante AddRelativeForce
    public float moveForce = 10.0f;

    // Dentro de ProcesarInput():
    if (Input.GetKey(KeyCode.W))
    {
        rb.AddRelativeForce(Vector3.forward * moveForce);
    }
    else if (Input.GetKey(KeyCode.S))
    {
        rb.AddRelativeForce(Vector3.back * moveForce);
    }"""
    add_cell_code_block(cell_ej, code_step4_2, cfg)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig9_code_addrelativeforce.png"),
        "Figura 9: Sustitución de rotación por fuerzas dinámicas relativas (rb.AddRelativeForce) para el avance y retroceso.",
        cfg)

    # =========================================================================
    # ETAPA 5
    # =========================================================================
    add_cell_heading_2(cell_ej, "Etapa 5: Diagnóstico del Eje Frontal Invertido del Modelo y Validación Consolidada", cfg)
    add_cell_body_paragraph(cell_ej, 
        "En esta etapa final se diagnostica y soluciona la desalineación de los ejes locales del modelo 3D y se consolida el controlador final del personaje.",
        cfg
    )

    # Subpaso 5.1
    add_cell_bullet_point(cell_ej, "Paso 5.1: Inspección del Gizmo de Coordenadas Locales (Z-Forward Invertido). ", 
        "Al probar en Unity el código con 'AddRelativeForce(Vector3.forward * moveForce)', se detectó un comportamiento anómalo: al presionar 'W', el personaje "
        "se desplazaba hacia atrás, y al presionar 'S', avanzaba. Al inspeccionar el GameObject 'Mango' en la vista Scene con la herramienta Gizmo en modo 'Local', "
        "se comprobó la causa raíz: la flecha azul (eje local Z positivo, 'forward') apuntaba directamente hacia la espalda y cola del gato. Esta desalineación "
        "es común en assets tridimensionales importados cuyos pivotes fueron exportados con una orientación invertida respecto a la malla visible.", 
        cfg, level=1)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig10_z_axis_gizmo.png"),
        "Figura 10: Vista Scene evidenciando que el eje frontal local Z (flecha azul) apunta hacia la espalda del modelo 3D.",
        cfg)

    # Subpaso 5.2
    add_cell_bullet_point(cell_ej, "Paso 5.2: Inversión de Vectores de Control en Código Fuente. ", 
        "Para corregir el avance sin tener que reexportar la malla 3D ni desajustar el colisionador, se implementa una solución defensiva en el código C#: "
        "se asigna 'Vector3.back' a la tecla 'W' (empujando hacia la dirección donde visualmente mira el gato) y 'Vector3.forward' a la tecla 'S' (retroceso). "
        "Asimismo, se ajusta la magnitud de 'moveForce' a un valor adecuado para evitar desplazamientos excesivamente violentos:", 
        cfg, level=1)

    code_step5_2 = """// CatControl.cs - Fragmento Incremental: Corrección de Ejes Invertidos
    if (Input.GetKey(KeyCode.W))
    {
        rb.AddRelativeForce(Vector3.back * moveForce);
    }
    else if (Input.GetKey(KeyCode.S))
    {
        rb.AddRelativeForce(Vector3.forward * moveForce);
    }"""
    add_cell_code_block(cell_ej, code_step5_2, cfg)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig11_code_final_inverted.png"),
        "Figura 11: Mapeo defensivo corregido en C#: 'W' asociado a Vector3.back y 'S' a Vector3.forward.",
        cfg)

    # Subpaso 5.3
    add_cell_bullet_point(cell_ej, "Paso 5.3: Demostración Final del Controlador Físico Consolidado. ", 
        "Con todas las correcciones integradas, se ejecuta la escena final en modo Play. El controlador de 'Mango' opera con perfecta sincronía: avanza hacia su frente "
        "con 'W', retrocede con 'S', rota con suavidad cinematográfica independiente de cuadros con 'A'/'D', y salta limpiamente con la barra espaciadora al aplicar "
        "fuerza relativa en Y, interactuando de forma robusta con la plataforma y el terreno de la escena.", 
        cfg, level=1)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig12_final_movement.gif"),
        "Figura 12 (GIF Animado): Validación final del ciclo de locomoción completo: avance, retroceso, giro suave y salto físico.",
        cfg)

    # Subpaso 5.4
    add_cell_bullet_point(cell_ej, "Paso 5.4: Arquitectura de Componentes en el Inspector y Código Consolidado. ", 
        "En el Inspector de Unity, el GameObject 'Mango' integra un Rigidbody (con masa unitaria, gravedad activa y restricciones 'Freeze Rotation' en X, Y, Z activadas "
        "para evitar torque parásito por fricción), un CapsuleCollider ajustado a la anatomía del gato y el componente 'CatControl' con los parámetros calibrados "
        "(Thrust = 5, Rotation Speed = 50, Move Force = 3-10):", 
        cfg, level=1)

    add_cell_image_with_caption(cell_ej,
        os.path.join(media_dir, "fig13_inspector_setup.png"),
        "Figura 13: Configuración integral del Inspector de Unity mostrando Rigidbody, CapsuleCollider y CatControl.",
        cfg)

    add_cell_body_paragraph(cell_ej, "A continuación se presenta el código fuente completo, consolidado y documentado listo para producción:", cfg)

    full_code = """using UnityEngine;

public class CatControl : MonoBehaviour
{
    [Header("Parámetros de Movimiento y Salto")]
    public float thrust = 5.0f;          // Fuerza vertical de salto
    public float rotationSpeed = 50.0f;  // Velocidad angular en grados/segundo
    public float moveForce = 3.0f;       // Magnitud de empuje traslacional relativo

    Rigidbody rb;

    void Start()
    {
        // Obtención de la referencia al componente Rigidbody nativo de la escena
        rb = GetComponent<Rigidbody>();
    }

    void Update()
    {
        // Lectura continua y procesamiento de eventos de teclado en cada frame
        ProcesarInput();
    }

    private void ProcesarInput()
    {
        // 1. Salto Vertical mediante fuerza física relativa
        if (Input.GetKey(KeyCode.Space))
        {
            rb.AddRelativeForce(Vector3.up * thrust);
        }

        // 2. Rotación Angular sobre el Eje Y (Independiente de la tasa de FPS)
        if (Input.GetKey(KeyCode.D))
        {
            // Giro horario (derecha) escalado por Time.deltaTime
            transform.Rotate(Vector3.up * rotationSpeed * Time.deltaTime);
        }
        else if (Input.GetKey(KeyCode.A))
        {
            // Giro antihorario (izquierda) escalado por Time.deltaTime
            transform.Rotate(Vector3.down * rotationSpeed * Time.deltaTime);
        }

        // 3. Traslación Frontal Relativa (Compensando eje Z invertido del modelo)
        if (Input.GetKey(KeyCode.W))
        {
            // Eje frontal visual del modelo apunta hacia Vector3.back
            rb.AddRelativeForce(Vector3.back * moveForce);
        }
        else if (Input.GetKey(KeyCode.S))
        {
            // Retroceso relativo sobre el eje Z local
            rb.AddRelativeForce(Vector3.forward * moveForce);
        }
    }
}"""
    add_cell_code_block(cell_ej, full_code, cfg)

    # --- 3. SECCIÓN: CUESTIONARIO TÉCNICO ---
    cell_cue = table.rows[10].cells[0]
    cell_cue.text = ""
    apply_cell_margins(cell_cue)

    add_cell_heading_1(cell_cue, "CUESTIONARIO TÉCNICO", cfg)

    # Pregunta 1
    add_cell_body_paragraph(cell_cue, 
        "¿A qué se le conoce como “Independencia de Frames” en Unity?",
        cfg, bold_prefix="Pregunta 1: "
    )
    add_cell_body_paragraph(cell_cue, 
        "En el desarrollo de videojuegos y motores gráficos como Unity, se denomina 'Independencia de Frames' (Frame Rate Independence) "
        "al principio y diseño arquitectónico según el cual las simulaciones temporales (traslaciones, rotaciones angulares, temporizadores, "
        "enfriamientos de habilidades y animaciones procedimentales) progresan a una velocidad constante y homogénea en tiempo real (por segundo), "
        "con total independencia de la cantidad de fotogramas que la CPU/GPU del cliente sea capaz de renderizar por segundo (FPS).",
        cfg
    )
    add_cell_body_paragraph(cell_cue, 
        "En una arquitectura dependiente de frames, cualquier operación que aplique un cambio fijo dentro de Update() (por ejemplo, sumar 50 grados "
        "en cada ciclo) producirá un resultado proporcional al hardware: en una computadora de alta gama que ejecute la escena a 240 FPS, el objeto "
        "rotará 240 veces por segundo (12,000 grados/segundo, provocando el fenómeno de 'objeto poseído' verificado en el laboratorio); mientras que "
        "en un dispositivo móvil o gama baja a 30 FPS rotará únicamente a 1,500 grados/segundo. La independencia de frames desacopla el ciclo de "
        "renderizado del ciclo de simulación temporal, garantizando paridad física, justicia competitiva y consistencia cinemática universal entre "
        "dispositivos heterogéneos.",
        cfg
    )

    # Pregunta 2
    add_cell_body_paragraph(cell_cue, 
        "¿Qué función realiza la propiedad Time.DeltaTime?",
        cfg, bold_prefix="Pregunta 2: "
    )
    add_cell_body_paragraph(cell_cue, 
        "La propiedad 'Time.deltaTime' es un atributo estático de solo lectura perteneciente a la clase UnityEngine.Time que devuelve un número "
        "en coma flotante (float) que representa el intervalo de tiempo exacto, medido en segundos, transcurrido entre la finalización del fotograma anterior "
        "y el inicio del fotograma en curso.",
        cfg
    )
    add_cell_body_paragraph(cell_cue, 
        "Matemáticamente, actúa como un factor de conversión diferencial entre una tasa expresada 'por segundo' y el incremento infinitesimal aplicable "
        "'por fotograma'. Por ejemplo, si el juego se ejecuta a 60 FPS, Time.deltaTime tiene un valor medio de 1/60 ≈ 0.01667 segundos; si corre a 240 FPS, "
        "su valor desciende a 1/240 ≈ 0.00417 segundos. Mediante el análisis dimensional:",
        cfg
    )
    add_cell_body_paragraph(cell_cue, 
        "Velocidad Angular (grados / segundo) × Time.deltaTime (segundos / frame) = Desplazamiento Angular (grados / frame)",
        cfg, italic=True
    )
    add_cell_body_paragraph(cell_cue, 
        "Al integrar esta suma infinitesimal en Update() a lo largo de un segundo exacto, la suma de los valores de Time.deltaTime de todos los cuadros "
        "completados converge matemáticamente a 1.0 segundo (Σ Time.deltaTime = 1.0 s). De este modo, la rotación o desplazamiento total acumulado "
        "en un segundo coincide exactamente con el valor nominal de la variable (ej. 50 grados reales), amortiguando de forma automática cualquier caída, "
        "pico o fluctuación de rendimiento en el hardware.",
        cfg
    )

    # Pregunta 3
    add_cell_body_paragraph(cell_cue, 
        "¿Cómo podemos hacer uso del componente Transform a través de código sin utilizar GetComponent?",
        cfg, bold_prefix="Pregunta 3: "
    )
    add_cell_body_paragraph(cell_cue, 
        "En Unity, se puede hacer uso del componente Transform de forma directa y nativa a través de la propiedad pública heredada 'transform' "
        "(con inicial minúscula), disponible automáticamente en cualquier script que derive de la clase base MonoBehaviour (y a su vez de Component):",
        cfg
    )
    add_cell_body_paragraph(cell_cue, 
        "public Transform transform { get; }",
        cfg, italic=True
    )
    add_cell_body_paragraph(cell_cue, 
        "Desde la evolución de las APIs internas de Unity (Unity 5 en adelante), esta propiedad no es una llamada sintáctica diferida a GetComponent<Transform>(), "
        "sino un acceso directo y optimizado por puntero a nivel de C++ nativo dentro del motor. Todo GameObject en Unity cuenta de manera obligatoria y ontológica "
        "con una raíz espacial representada por su Transform (no existe en Unity un GameObject que carezca de Transform), lo que permite que el motor garantice "
        "que la referencia jamás sea nula y esté preenlazada en memoria con complejidad de acceso O(1).",
        cfg
    )
    add_cell_body_paragraph(cell_cue, 
        "En consecuencia, el programador puede manipular directamente la orientación y posición escribiendo sentencias como "
        "'transform.Rotate(...)', 'transform.Translate(...)', 'transform.position', 'transform.rotation' o 'transform.forward' en cualquier método de la clase, "
        "ahorrando la sobrecarga computacional de búsqueda en tablas de componentes y haciendo el código mucho más conciso y eficiente.",
        cfg
    )

    # --- 4. SECCIÓN: CONCLUSIONES ---
    cell_concl = table.rows[13].cells[0]
    cell_concl.text = ""
    apply_cell_margins(cell_concl)

    add_cell_heading_1(cell_concl, "CONCLUSIONES", cfg)

    add_cell_bullet_point(cell_concl, 
        "Garantía de Consistencia Mediante Time.deltaTime: ", 
        "Se demostró experimentalmente que omitir el escalado por tiempo en el ciclo Update() genera una ruptura crítica en la física del juego "
        "a altas tasas de refresco (230+ FPS), donde el personaje rota a velocidades extremas e incontrolables. La inclusión de Time.deltaTime es el estándar "
        "indispensable para garantizar que las transformaciones espaciales progresen a velocidad constante por segundo, desacopladas de las capacidades del hardware.", 
        cfg, level=1)

    add_cell_bullet_point(cell_concl, 
        "Diferenciación entre Cinemática Angular y Dinámica de Fuerzas: ", 
        "El experimento reveló la incompatibilidad de reutilizar 'transform.Rotate' para traslaciones longitudinales (W/S), lo cual provocaba cabeceo "
        "y volteretas anómalas sobre el eje transversal. La aplicación de 'rb.AddRelativeForce' sobre el Rigidbody demostró ser la metodología óptima "
        "para desplazar al cuerpo rígido respetando su sistema de coordenadas local y preservando la simulación de colisiones del motor PhysX.", 
        cfg, level=1)

    add_cell_bullet_point(cell_concl, 
        "Resolución Defensiva de Inversión de Ejes en Mallas 3D: ", 
        "Se constató que los modelos importados pueden presentar discrepancias entre el eje frontal visual de la malla y el vector local 'Vector3.forward' "
        "(eje Z azul apuntando hacia la cola). La adaptación en código mediante la inversión de los vectores de fuerza ('Vector3.back' para avance) "
        "constituye una solución programática elegante y eficiente que evita la reexportación destructiva del asset en programas de modelado externos.", 
        cfg, level=1)

    add_cell_bullet_point(cell_concl, 
        "Eficiencia en el Acceso Directo a Componentes Nativos: ", 
        "El aprovechamiento de la propiedad predefinida 'transform' en scripts derivados de MonoBehaviour optimiza significativamente la legibilidad "
        "y el rendimiento de ejecución, eliminando búsquedas innecesarias mediante 'GetComponent<Transform>()' al operar con punteros cacheados "
        "directamente en la capa nativa del motor de juego.", 
        cfg, level=1)

    # --- 5. SECCIÓN: REFERENCIAS Y BIBLIOGRAFÍA ---
    cell_ref = table.rows[17].cells[0]
    cell_ref.text = ""
    apply_cell_margins(cell_ref)

    add_cell_heading_1(cell_ref, "REFERENCIAS Y BIBLIOGRAFÍA", cfg)

    add_cell_bullet_point(cell_ref, 
        "Unity Technologies. (2024). ", 
        "Time.deltaTime - Unity Scripting API Reference (v6000.0). San Francisco, CA: Unity Documentation. Recuperado de https://docs.unity.com/ScriptReference/Time-deltaTime.html", 
        cfg, level=1)

    add_cell_bullet_point(cell_ref, 
        "Unity Technologies. (2024). ", 
        "Transform.Rotate - Unity Scripting API Reference. San Francisco, CA: Unity Documentation. Recuperado de https://docs.unity.com/ScriptReference/Transform.Rotate.html", 
        cfg, level=1)

    add_cell_bullet_point(cell_ref, 
        "Unity Technologies. (2024). ", 
        "Rigidbody.AddRelativeForce - Unity Scripting API Reference. San Francisco, CA: Unity Documentation. Recuperado de https://docs.unity.com/ScriptReference/Rigidbody.AddRelativeForce.html", 
        cfg, level=1)

    add_cell_bullet_point(cell_ref, 
        "Gregory, J. (2018). ", 
        "Game Engine Architecture (3.ª ed.). Boca Raton, FL: CRC Press / Taylor & Francis Group.", 
        cfg, level=1)

    add_cell_bullet_point(cell_ref, 
        "Nystrom, R. (2014). ", 
        "Game Programming Patterns: Game Loop and Update Method Patterns. Genever Benning.", 
        cfg, level=1)

    # Guardar documento
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"✓ Reporte compilado exitosamente en: {output_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generador de reporte de Laboratorio 06")
    parser.add_argument("--template", default="inputs/IDSE/IDSE-Lab05/input-template.docx")
    parser.add_argument("--output", default="output/IDSE/lab06/Reporte_Practica_06_Unity.docx")
    parser.add_argument("--media", default="output/IDSE/lab06/figures")
    args = parser.parse_args()

    build_lab06_report(
        template_path=args.template,
        output_path=args.output,
        media_dir=args.media
    )
