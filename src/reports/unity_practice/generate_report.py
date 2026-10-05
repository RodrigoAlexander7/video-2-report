import os
import argparse
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from config import ReportConfig, DEFAULT_CONFIG
from src.video_engine.media_extractor import generate_and_enhance_figures

# Paleta corporativa sobria
COLOR_PRIMARY_NAVY = RGBColor(0x1B, 0x36, 0x5D)    # Azul institucional #1B365D
COLOR_SECONDARY_DARK = RGBColor(0x2B, 0x4C, 0x7E)  # Azul medio #2B4C7E
COLOR_BODY_TEXT = RGBColor(0x1F, 0x24, 0x21)       # Grafito suave para lectura
COLOR_CAPTION = RGBColor(0x55, 0x5B, 0x6E)         # Gris tenue
COLOR_BORDER = "CBD5E1"                            # Borde tenue para bloques de código
COLOR_CODE_BG = "F8FAFC"                           # Fondo bloque código #F8FAFC
COLOR_NOTE_BG = "EFF6FF"                           # Fondo caja informativa #EFF6FF

def apply_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Establece márgenes internos de respiración en las celdas de la tabla (en dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_1(cell, text, cfg: ReportConfig):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = cfg.font_family_body
    run.font.size = Pt(cfg.font_size_heading1_pt)
    run.font.color.rgb = COLOR_PRIMARY_NAVY
    return p

def add_heading_2(cell, text, cfg: ReportConfig):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = cfg.font_family_body
    run.font.size = Pt(cfg.font_size_heading2_pt)
    run.font.color.rgb = COLOR_SECONDARY_DARK
    return p

def add_body_paragraph(cell, text, cfg: ReportConfig, bold_prefix="", italic=False):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(cfg.paragraph_space_after_pt)
    p.paragraph_format.line_spacing = cfg.body_line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = cfg.font_family_body
        r_pre.font.size = Pt(cfg.font_size_body_pt)
        r_pre.font.color.rgb = COLOR_BODY_TEXT
    
    run = p.add_run(text)
    run.font.name = cfg.font_family_body
    run.font.size = Pt(cfg.font_size_body_pt)
    run.font.color.rgb = COLOR_BODY_TEXT
    run.italic = italic
    return p

def add_bullet_point(cell, bold_text, desc_text, cfg: ReportConfig, level: int = 1):
    p = cell.add_paragraph()
    indent = cfg.bullet_indent_inches if level == 1 else cfg.bullet_sub_indent_inches
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = cfg.body_line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    bullet_symbol = "■  " if level == 1 else "○  "
    r_sym = p.add_run(bullet_symbol)
    r_sym.font.name = cfg.font_family_body
    r_sym.font.size = Pt(cfg.font_size_body_pt - 1)
    r_sym.font.color.rgb = COLOR_PRIMARY_NAVY if level == 1 else COLOR_SECONDARY_DARK

    if bold_text:
        r_b = p.add_run(bold_text)
        r_b.bold = True
        r_b.font.name = cfg.font_family_body
        r_b.font.size = Pt(cfg.font_size_body_pt)
        r_b.font.color.rgb = COLOR_BODY_TEXT

    r_d = p.add_run(desc_text)
    r_d.font.name = cfg.font_family_body
    r_d.font.size = Pt(cfg.font_size_body_pt)
    r_d.font.color.rgb = COLOR_BODY_TEXT
    return p

def add_code_block(cell, code_str, cfg: ReportConfig):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    
    pPr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{COLOR_CODE_BG}"/>')
    pPr.append(shd)

    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="18" w:space="8" w:color="1B365D"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="4" w:color="{COLOR_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="4" w:color="{COLOR_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="4" w:color="{COLOR_BORDER}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

    run = p.add_run(code_str)
    run.font.name = cfg.font_family_code
    run.font.size = Pt(cfg.font_size_code_pt)
    run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
    return p

def add_image_with_caption(cell, img_path, caption_text, cfg: ReportConfig, width=Inches(5.4)):
    p_img = cell.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.add_run().add_picture(img_path, width=width)
    
    p_cap = cell.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = cfg.font_family_body
    r_cap.font.size = Pt(cfg.font_size_caption_pt)
    r_cap.italic = True
    r_cap.font.color.rgb = COLOR_CAPTION

def build_report(cfg: ReportConfig = DEFAULT_CONFIG):
    doc = docx.Document(cfg.template_path)
    table = doc.tables[0]

    # --- 1. ENCABEZADO Y TÍTULO ---
    table.rows[2].cells[1].text = "Implementación del Sistema de Movimiento Físico y Salto para Personaje en Unity 3D"
    p_tit = table.rows[2].cells[1].paragraphs[0]
    p_tit.runs[0].font.name = cfg.font_family_body
    p_tit.runs[0].font.size = Pt(cfg.font_size_body_pt)
    p_tit.runs[0].bold = True

    # --- 2. SECCIÓN: RESULTADOS Y PRUEBAS (TUTORIAL PROGRESIVO) ---
    cell_ej = table.rows[9].cells[0]
    cell_ej.text = ""
    apply_cell_margins(cell_ej)

    add_heading_1(cell_ej, "EJERCICIOS RESUELTOS: TUTORIAL DE CONTROLADOR DE PERSONAJE (CatControl)", cfg)

    add_body_paragraph(cell_ej, 
        "A continuación se presenta la guía instructiva y metodológica desarrollada a lo largo de la práctica para implementar "
        "desde cero el sistema de movimiento físico y salto de un personaje tridimensional ('Mango') en Unity 6. "
        "El tutorial está organizado en 5 etapas secuenciales, abordando desde la configuración de componentes hasta la resolución "
        "de problemas físicos reales observados en ejecución.",
        cfg
    )

    # --- ETAPA 1 ---
    add_heading_2(cell_ej, "Etapa 1: Preparación del GameObject y Configuración de Físicas en Unity", cfg)
    add_body_paragraph(cell_ej, 
        "En esta primera etapa se dota al modelo tridimensional de la capacidad de interactuar con la gravedad y colisionar con el suelo.",
        cfg
    )
    add_bullet_point(cell_ej, "Paso 1.1: Seleccionar el GameObject. ", "En la jerarquía de la escena, localice y seleccione el objeto del personaje ('Mango').", cfg, level=1)
    add_bullet_point(cell_ej, "Paso 1.2: Añadir componentes físicos. ", "En el Inspector, agregue un componente 'Rigidbody' (con masa unitaria y 'Use Gravity' habilitado) y un 'CapsuleCollider' ajustando su radio y altura a la malla del modelo.", cfg, level=1)
    add_bullet_point(cell_ej, "Paso 1.3: Asignar el script. ", "Cree y vincule el script C# denominado CatControl.cs, garantizando la directiva [RequireComponent(typeof(Rigidbody))] para forzar la presencia del motor físico.", cfg, level=1)

    add_image_with_caption(cell_ej, 
        os.path.join(cfg.media_dir, "unity_inspector_rigidbody.png"),
        "Figura 1: Configuración de componentes físicos en el Inspector de Unity (destacando las restricciones requeridas).",
        cfg)

    # --- ETAPA 2 ---
    add_heading_2(cell_ej, "Etapa 2: Declaración de Parámetros y Lectura de Entradas (Movimiento X/Z)", cfg)
    add_body_paragraph(cell_ej, 
        "El objetivo de esta etapa es capturar las pulsaciones de teclado del usuario (WASD) y traducirlas a un vector de dirección tridimensional normalizado, "
        "aplicando una velocidad tangible sobre el plano horizontal sin cancelar la velocidad vertical existente en Y.",
        cfg
    )
    add_bullet_point(cell_ej, "Instrucción de codificación: ", "Declare las variables públicas de velocidad en la cabecera de la clase y defina el método 'ProcessInput()' invocado desde 'Update()':", cfg, level=1)

    code_step2 = """// CatControl.cs - Fragmento Incremental: Variables y Movimiento
[RequireComponent(typeof(Rigidbody))]
public class CatControl : MonoBehaviour
{
    public float speed = 5.0f;
    public float jumpForce = 5.0f;
    Rigidbody rb;

    void Start()
    {
        rb = GetComponent<Rigidbody>();
    }

    void Update()
    {
        ProcessInput();
    }

    private void ProcessInput()
    {
        float x = 0f;
        float z = 0f;

        if (Input.GetKey(KeyCode.W)) z += 1f;
        else if (Input.GetKey(KeyCode.S)) z -= 1f;
        if (Input.GetKey(KeyCode.A)) x -= 1f;
        else if (Input.GetKey(KeyCode.D)) x += 1f;

        // Normalización para velocidad equitativa en diagonales
        Vector3 direction = new Vector3(x, 0f, z).normalized;

        // Asignación directa a linearVelocity respetando la caída gravitacional previa
        rb.linearVelocity = new Vector3(direction.x * speed, rb.linearVelocity.y, direction.z * speed);
    }
}"""
    add_code_block(cell_ej, code_step2, cfg)

    # --- ETAPA 3 ---
    add_heading_2(cell_ej, "Etapa 3: Prueba de Desplazamiento y Corrección de Torque Parásito", cfg)
    add_body_paragraph(cell_ej, 
        "Una vez programado el movimiento horizontal, se procede a ingresar en modo Play en Unity para evaluar la respuesta del personaje sobre el terreno.",
        cfg
    )
    add_bullet_point(cell_ej, "Incidencia Observada (Minuto 13:38): ", 
        "Al presionar las teclas de movimiento, la fricción producida en la base del colisionador genera un momento de fuerza (torque) "
        "que hace que el gato pierda el equilibrio, tambaleándose y volcando descontroladamente sobre sus ejes horizontales.", cfg, level=1)

    add_image_with_caption(cell_ej, 
        os.path.join(cfg.media_dir, "animacion_movimiento_rotacion_error.gif"),
        "Figura 2 (GIF Animado): Anomalía de torque y vuelco postural por falta de restricciones angulares.",
        cfg)

    add_bullet_point(cell_ej, "Instrucción Correctiva en Unity: ", 
        "Para solucionar este vuelco físico, regrese al Inspector del objeto 'Mango' -> Rigidbody -> Constraints, y marque las casillas "
        "'Freeze Rotation' en X, Y y Z (señaladas por la flecha en la Figura 1). Con esto, el motor físico ignora las fuerzas de rotación parásitas.", cfg, level=1)

    add_image_with_caption(cell_ej, 
        os.path.join(cfg.media_dir, "animacion_movimiento_exitoso.gif"),
        "Figura 3 (GIF Animado): Desplazamiento rectilíneo fluido y equilibrado tras activar Freeze Rotation en el Inspector.",
        cfg)

    # --- ETAPA 4 ---
    add_heading_2(cell_ej, "Etapa 4: Implementación de la Detección de Suelo y Mecánica de Salto", cfg)
    add_body_paragraph(cell_ej, 
        "Para incorporar el salto sin permitir que el jugador ejecute saltos infinitos en el aire (Double Jump indeseado), "
        "se introduce una bandera booleana 'isGrounded' condicionada a los eventos físicos de contacto contra el suelo.",
        cfg
    )
    add_bullet_point(cell_ej, "Paso 4.1: Condición de Salto. ", "Al presionar la barra espaciadora (KeyCode.Space), se inyecta el valor de 'jumpForce' en el eje Y de 'linearVelocity', solo si 'isGrounded' es verdadero.", cfg, level=1)
    add_bullet_point(cell_ej, "Paso 4.2: Callbacks Mágicos de MonoBehaviour. ", "Se emplean 'OnCollisionStay' para mantener activa la condición de suelo mientras se esté en contacto con el terreno, y 'OnCollisionExit' para revocarla de inmediato al despegarse.", cfg, level=1)

    code_step4 = """// CatControl.cs - Fragmento Incremental: Lógica de Salto y Colisiones
    bool isGrounded;

    private void ProcessInput()
    {
        // ... (código de movimiento previo)

        if (Input.GetKeyDown(KeyCode.Space) && isGrounded)
        {
            rb.linearVelocity = new Vector3(rb.linearVelocity.x, jumpForce, rb.linearVelocity.z);
        }
    }

    // Callbacks situados en el ámbito DIRECTO de la clase MonoBehaviour:
    void OnCollisionStay(Collision collision)
    {
        isGrounded = true;
    }

    void OnCollisionExit(Collision collision)
    {
        isGrounded = false;
    }"""
    add_code_block(cell_ej, code_step4, cfg)

    add_bullet_point(cell_ej, "Advertencia Técnica Clave (Ámbito de Clases): ", 
        "Durante la sesión se cometió el error inicial de anidar 'OnCollisionStay' dentro del método 'ProcessInput()'. "
        "Debido a que Unity invoca estos eventos mágicos por reflexión buscando la firma en la raíz de MonoBehaviour, nunca se ejecutaban y el salto fallaba. "
        "Asegúrese rigurosamente de que ambos métodos se ubiquen fuera de 'ProcessInput()' y al mismo nivel jerárquico que 'Start()' y 'Update()'.", cfg, level=1)

    # --- ETAPA 5 ---
    add_heading_2(cell_ej, "Etapa 5: Verificación Final y Código Fuente Consolidado", cfg)
    add_body_paragraph(cell_ej, 
        "Tras corregir el anidamiento de los métodos mágicos, se ejecuta la escena final en Unity para constatar la sinergia "
        "entre el desplazamiento multidireccional y la mecánica de salto continuo.",
        cfg
    )

    add_image_with_caption(cell_ej, 
        os.path.join(cfg.media_dir, "animacion_salto_gato.gif"),
        "Figura 4 (GIF Animado): Validación final del salto continuo responsivo respetando las colisiones del terreno.",
        cfg)

    add_body_paragraph(cell_ej, "Como referencia completa para el lector, a continuación se consolida el script íntegro listo para producción:", cfg)

    full_code = """using UnityEngine;

[RequireComponent(typeof(Rigidbody))]
public class CatControl : MonoBehaviour
{
    public float speed = 5.0f;
    public float jumpForce = 5.0f;

    Rigidbody rb;
    bool isGrounded;

    void Start()
    {
        rb = GetComponent<Rigidbody>();
    }

    void Update()
    {
        ProcessInput();
    }

    private void ProcessInput()
    {
        float x = 0f;
        float z = 0f;

        if (Input.GetKey(KeyCode.W)) z += 1f;
        else if (Input.GetKey(KeyCode.S)) z -= 1f;
        if (Input.GetKey(KeyCode.A)) x -= 1f;
        else if (Input.GetKey(KeyCode.D)) x += 1f;

        Vector3 direction = new Vector3(x, 0f, z).normalized;

        rb.linearVelocity = new Vector3(
            direction.x * speed,
            rb.linearVelocity.y,
            direction.z * speed
        );

        if (Input.GetKeyDown(KeyCode.Space) && isGrounded)
        {
            rb.linearVelocity = new Vector3(rb.linearVelocity.x, jumpForce, rb.linearVelocity.z);
        }
    }

    void OnCollisionStay(Collision collision)
    {
        isGrounded = true;
    }

    void OnCollisionExit(Collision collision)
    {
        isGrounded = false;
    }
}"""
    add_code_block(cell_ej, full_code, cfg)

    # --- 3. SECCIÓN: CUESTIONARIO TÉCNICO ---
    cell_cue = table.rows[10].cells[0]
    cell_cue.text = ""
    apply_cell_margins(cell_cue)

    add_heading_1(cell_cue, "CUESTIONARIO TÉCNICO", cfg)

    add_body_paragraph(cell_cue, 
        "¿Cuál es el rol de la propiedad linearVelocity en Unity 6 y por qué se prefiere frente a modificar transform.position o aplicar AddForce en este controlador?",
        cfg, bold_prefix="Pregunta 1: "
    )
    add_body_paragraph(cell_cue, 
        "En la arquitectura de físicas de Unity 6, 'linearVelocity' sustituye a la propiedad obsoleta 'velocity' del componente Rigidbody. "
        "A diferencia de modificar 'transform.position' (que altera las coordenadas cartesianas de forma discreta 'teletransportando' al objeto y "
        "vulnerando el cálculo de colisiones continuas en el motor PhysX), 'linearVelocity' asigna un vector de desplazamiento instantáneo respetando "
        "la integración del solver. Frente a 'AddForce' (que añade aceleración newtoniana sujeta a acumulación e inercia), 'linearVelocity' ofrece un control "
        "arcade de respuesta inmediata, deteniendo o acelerando al personaje en el frame exacto de pulsación del teclado sin deslizamientos indeseados.",
        cfg
    )

    add_body_paragraph(cell_cue, 
        "¿Por qué es indispensable la normalización vectorial (.normalized) al construir el vector de dirección?",
        cfg, bold_prefix="Pregunta 2: "
    )
    add_body_paragraph(cell_cue, 
        "Al presionar combinaciones de teclas simultáneas (por ejemplo, W y D para avanzar en diagonal), la suma de componentes produce un vector (1, 0, 1). "
        "Su norma escalar es √(1² + 0² + 1²) = √2 ≈ 1.4142. Si no se normalizara, la velocidad lineal en diagonales superaría en un 41.4% a la velocidad en trayectorias "
        "axiales puras. Al aplicar el operador '.normalized', el vector preserva su rumbo pero fija su magnitud exactamente en 1, garantizando paridad física uniforme.",
        cfg
    )

    add_body_paragraph(cell_cue, 
        "¿Qué función desempeñan OnCollisionStay y OnCollisionExit en el control de salto y qué recaudos técnicos deben considerarse en escenarios complejos?",
        cfg, bold_prefix="Pregunta 3: "
    )
    add_body_paragraph(cell_cue, 
        "Ambos callbacks gestionan el estado booleano 'isGrounded'. 'OnCollisionStay' valida en cada ciclo de la simulación física (FixedUpdate) que los colisionadores "
        "continúan en contacto, habilitando el salto. 'OnCollisionExit' revoca la condición en cuanto el cuerpo se despega del plano, evitando el salto infinito "
        "(Double Jump no programado). En producciones avanzadas, se recomienda verificar que el vector normal del punto de colisión apunte hacia arriba (contact.normal.y > 0.7), "
        "para evitar que rozar una pared vertical reactive erróneamente la condición de suelo.",
        cfg
    )

    # --- 4. SECCIÓN: CONCLUSIONES ---
    cell_concl = table.rows[13].cells[0]
    cell_concl.text = ""
    apply_cell_margins(cell_concl)

    add_heading_1(cell_concl, "CONCLUSIONES", cfg)

    add_bullet_point(cell_concl, 
        "Control Físico Cinemático Híbrido: ", 
        "La manipulación directa de linearVelocity en los componentes horizontales (X, Z) combinada con la preservación del componente vertical preexistente (rb.linearVelocity.y) "
        "constituye un paradigma altamente eficiente para videojuegos de acción, pues otorga respuesta inmediata a la entrada del usuario sin interferir con la aceleración gravitacional natural.",
        cfg, level=1)

    add_bullet_point(cell_concl, 
        "Trascendencia de las Restricciones Angulares (Constraints): ", 
        "Quedó demostrado experimentalmente que un cuerpo rígido tridimensional sometido a fricciones de traslación experimenta momentos de vuelco inevitables. "
        "El bloqueo explícito de rotaciones angulares en el Rigidbody es un requisito indispensable para personajes que no utilicen modelos de muñeco de trapo (ragdolls).",
        cfg, level=1)

    add_bullet_point(cell_concl, 
        "Rigor en la Jerarquía de Mensajes de MonoBehaviour: ", 
        "La arquitectura basada en reflexión de Unity exige que los métodos de detección física se sitúen rigurosamente en el ámbito principal de la clase derivada. "
        "El encapsulamiento indebido de estas funciones bloquea la recepción de eventos del motor sin arrojar errores explícitos de compilación.",
        cfg, level=1)

    # --- 5. SECCIÓN: REFERENCIAS Y BIBLIOGRAFÍA ---
    cell_ref = table.rows[17].cells[0]
    cell_ref.text = ""
    apply_cell_margins(cell_ref)

    add_heading_1(cell_ref, "REFERENCIAS Y BIBLIOGRAFÍA", cfg)

    add_bullet_point(cell_ref, 
        "Unity Technologies. (2024). ", 
        "Rigidbody.linearVelocity - Unity Scripting API Reference (v6000.0). San Francisco, CA: Unity Documentation. Recuperado de https://docs.unity.com/ScriptReference/Rigidbody-linearVelocity.html",
        cfg, level=1)

    add_bullet_point(cell_ref, 
        "Unity Technologies. (2024). ", 
        "Order of Execution for Event Functions - Physics Lifecycle. Recuperado de https://docs.unity.com/Manual/ExecutionOrder.html",
        cfg, level=1)

    add_bullet_point(cell_ref, 
        "Nystrom, R. (2014). ", 
        "Game Programming Patterns: Component and Update Method Patterns. Genever Benning.",
        cfg, level=1)

    # Guardar documento
    doc.save(cfg.output_path)
    print(f"✓ Reporte compilado exitosamente en formato tutorial por etapas en: {cfg.output_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generador de reporte estilizado en Word")
    parser.add_argument("--low-quality", action="store_true", help="Generar GIFs en calidad estándar")
    args = parser.parse_args()

    cfg = ReportConfig(high_quality_gifs=not args.low_quality)
    build_report(cfg)
