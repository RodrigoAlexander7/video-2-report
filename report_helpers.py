#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador del Informe Académico Oficial de Negocios Electrónicos (Lab 05)
Tema: Gestión de la Cadena de Suministro (SCM/ERP) en Distribuidora Inca S.R.L.
Herramienta de Solución: Odoo Community Edition
Universidad Nacional de San Agustín de Arequipa (UNSA) - EPIS
Docente: Dr. Ing. César Basilio Baluarte Araya
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    """
    Configura bordes de celda.
    kwargs: top, bottom, left, right. Valores: 'single', 'double', 'dashed', 'none'
    color: hexadecimal sin '#', sz: grosor en octavos de punto
    """
    tcPr = cell._tc.get_or_add_tcPr()
    color = kwargs.get('color', 'CCCCCC')
    sz = kwargs.get('sz', '4')
    val = kwargs.get('val', 'single')
    
    top = kwargs.get('top', val)
    bottom = kwargs.get('bottom', val)
    left = kwargs.get('left', val)
    right = kwargs.get('right', val)
    
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22) # Tono sobrio carbón
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33) # Gris oscuro neutro
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return p

def add_p(doc, text, bold_prefix="", italic_prefix=False, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        if italic_prefix:
            r_pre.italic = True
    if text:
        r_txt = p.add_run(text)
        r_txt.font.name = 'Calibri'
        r_txt.font.size = Pt(11)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Agregar viñeta manual estilizada si no la toma el estilo
    r_bullet = p.add_run("▪  ")
    r_bullet.bold = True
    r_bullet.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
    if text:
        r_txt = p.add_run(text)
        r_txt.font.name = 'Calibri'
        r_txt.font.size = Pt(11)
    return p

def add_callout(doc, text, title=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F5F5F5")
    set_cell_border(cell, left='single', right='none', top='none', bottom='none', color='444444', sz='20')
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        rt = p.add_run(title + "\n")
        rt.bold = True
        rt.font.name = 'Calibri'
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    rtxt = p.add_run(text)
    rtxt.font.name = 'Calibri'
    rtxt.font.size = Pt(10.5)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(4)

def add_screenshot_placeholder(doc, fig_num, title, module, objective, test_data, visible_items, expected_result):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_shading(cell, "F8F9FA")
    set_cell_border(cell, left='single', right='single', top='single', bottom='single', color='444444', sz='10')
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(4)
    r_badge = p0.add_run(f"📷 EVIDENCIA DE PROTOTIPO - FIGURA {fig_num}: {title.upper()}\n")
    r_badge.bold = True
    r_badge.font.name = 'Calibri'
    r_badge.font.size = Pt(10.5)
    r_badge.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    
    r1 = p1.add_run("• Módulo Odoo: ")
    r1.bold = True
    p1.add_run(f"{module}\n")
    
    r2 = p1.add_run("• Objetivo de la prueba: ")
    r2.bold = True
    p1.add_run(f"{objective}\n")
    
    r3 = p1.add_run("• Datos de prueba simulados: ")
    r3.bold = True
    p1.add_run(f"{test_data}\n")
    
    r4 = p1.add_run("• Elementos visuales requeridos en la captura: ")
    r4.bold = True
    p1.add_run(f"{visible_items}\n")
    
    r5 = p1.add_run("• Resultado esperado: ")
    r5.bold = True
    p1.add_run(f"{expected_result}")
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(8)
    p2.paragraph_format.space_after = Pt(4)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_box = p2.add_run(
        "┌────────────────────────────────────────────────────────────────────────┐\n"
        "│                                                                        │\n"
        "│     [ PEGAR AQUÍ LA CAPTURA DE PANTALLA DE ODOO COMMUNITY UNA VEZ      │\n"
        "│                    EJECUTADA ESTA CONFIGURACIÓN ]                      │\n"
        "│                                                                        │\n"
        "└────────────────────────────────────────────────────────────────────────┘"
    )
    r_box.font.name = 'Consolas'
    r_box.font.size = Pt(8.5)
    r_box.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(6)
    rcap = p_cap.add_run(f"Figura {fig_num}. {title}. Fuente: Elaboración propia en Odoo Community Edition.")
    rcap.font.size = Pt(9.5)
    rcap.italic = True
    rcap.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

def add_table_custom(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Cabecera
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "333333") # Gris oscuro sobrio
        set_cell_border(hdr_cells[i], color="333333", sz="6")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Datos
    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F9FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_border(row_cells[c_idx], color="D3D3D3", sz="4")
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9)
                
    # Anchos
    if col_widths:
        for row in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)
                
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(4)

print("Helper functions defined successfully!")
