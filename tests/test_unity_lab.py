# -*- coding: utf-8 -*-
"""Reporte de laboratorio de Unity: modelo, estilo, tamaños, división de código, plantilla y flujo completo con un video sintético."""

import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import docx
from PIL import Image

from src.core.chrome import BrowserError, find_chrome
from src.editor_engine.extract import SnippetError
from src.editor_engine.split import split_range, split_ranges
from src.unity_lab import ReportError, check_lab, load_lab, scaffold
from src.unity_lab.code import plan_code
from src.unity_lab.layout import MAX_H_IN, MAX_W_IN, fit_capture, fit_code, prepare_bitmap
from src.unity_lab.model import CodeSpec
from src.unity_lab.style import Speller, check_citations, lint_text

CS = "\n".join(
    ["using UnityEngine;", "", "public class Cat : MonoBehaviour", "{"]
    + [f"    int campo{i} = {i};" for i in range(1, 11)]
    + ["", "    void Start()", "    {"]
    + [f"        campo{i} += 1;" for i in range(1, 25)]
    + ["    }", "}"])


def has_chrome() -> bool:
    try:
        find_chrome()
        return True
    except BrowserError:
        return False


def make_template(path: Path) -> Path:
    """Tabla con las filas de la plantilla EPIS (rótulos reales, celdas mínimas)."""
    d = docx.Document()
    t = d.add_table(rows=0, cols=6)

    def row(*cells, span=None):
        r = t.add_row()
        if span:
            merged = r.cells[0].merge(r.cells[5])
            merged.text = cells[0]
            return r
        for i, text in enumerate(cells):
            r.cells[i].text = text
        return r

    row("INFORMACIÓN BÁSICA", span=True)
    row("ASIGNATURA:", "CURSO", "", "", "", "")
    row("TÍTULO DE LA PRÁCTICA:", "Programación", "", "", "", "")
    row("NÚMERO DE PRÁCTICA:", "02", "AÑO LECTIVO:", "2026 - B", "NRO. SEMESTRE:", "Octavo")
    row("FECHA DE PRESENTACIÓN", "10/09/2026", "HORA DE PRESENTACIÓN", "12:00:00", "", "")
    r = row("INTEGRANTE (s)", "", "", "", "NOTA (0-20)", "")
    r.cells[0].add_paragraph("Nombre Viejo")
    r = row("DOCENTE(s):", span=True)
    r.cells[0].add_paragraph("Docente Viejo")
    row("RESULTADOS Y PRUEBAS", span=True)
    row("EJERCICIOS RESUELTOS:\nEjercicio 1 del enunciado", span=True)
    row("CUESTIONARIO:\n¿Describir?", span=True)
    row("CONCLUSIONES", span=True)
    row("Las prácticas realizadas", span=True)
    row("REFERENCIAS Y BIBLIOGRAFÍA", span=True)
    row("", span=True)
    d.save(path)
    return path


def write_lab(root: Path, report: str, video: bool = True, extra_meta: str = "") -> Path:
    (root / "content").mkdir(parents=True, exist_ok=True)
    (root / "code").mkdir(exist_ok=True)
    (root / "code/Cat.cs").write_text(CS, encoding="utf-8")
    make_template(root / "tpl.docx")
    (root / "project.toml").write_text(
        'pipeline = "unity-lab"\n[unity]\ntemplate = "tpl.docx"\n' + ('video = "v.mp4"\n' if video else "")
        + 'title = "Práctica de prueba"\nnumber = "09"\ndate = "01/01/2026"\nmembers = ["Ana Pérez", "Luis Gómez"]\n' + extra_meta,
        encoding="utf-8")
    (root / "content/report.toml").write_text(report, encoding="utf-8")
    return root


MINIMAL = '''
[[block]]
kind = "heading"
text = "Variables"

[[block]]
kind = "step"
text = "Declaramos los campos de la clase."
[block.code]
id = "campos"
file = "code/Cat.cs"
lines = "5-14"

[[question]]
q = "¿Qué hace Start?"
a = "Se ejecuta una vez al iniciar el objeto [1]."

[[conclusion]]
text = "Los campos quedaron declarados."

[[reference]]
text = "Unity Technologies, Scripting API, 2026."
'''


class ModelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def test_minimal_lab_loads(self):
        lab = load_lab(write_lab(self.dir, MINIMAL, video=False))
        self.assertEqual([b.kind for b in lab.blocks], ["heading", "step"])
        self.assertEqual(lab.blocks[1].code.id, "campos")

    def test_all_problems_are_reported_together(self):
        bad = '''
[[block]]
kind = "step"
text = "%s"
[block.figure]
id = "a"
caption = "Sin tiempo."
[[block]]
kind = "step"
text = "Otro."
[block.figure]
id = "a"
t = "2:xx"
caption = "Duplicada."
[[block]]
kind = "step"
text = "Falta archivo."
[block.code]
id = "c"
file = "code/NoExiste.cs"
[[block]]
kind = "banner"
''' % ("palabra " * 50)
        with self.assertRaises(ReportError) as cm:
            load_lab(write_lab(self.dir, bad))
        text = str(cm.exception)
        for fragment in ("palabras (máximo", "falta 't'", "ya se usa", "NoExiste.cs", "banner", "tiempo inválido"):
            self.assertIn(fragment, text)

    def test_meta_requires_the_basics(self):
        root = write_lab(self.dir, MINIMAL, video=False)
        (root / "project.toml").write_text('[unity]\ntitle = "x"\nfoo = 1\n', encoding="utf-8")
        with self.assertRaisesRegex(ReportError, "falta 'number'.*|foo"):
            load_lab(root)

    def test_pair_and_strip_specs(self):
        content = '''
[[block]]
kind = "step"
text = "Antes y después."
[block.figure]
id = "par"
kind = "pair"
caption = "Antes y después de la acción."
[[block.figure.frames]]
t = 10
[[block.figure.frames]]
t = 20
[[block]]
kind = "figure"
[block.figure]
id = "tira"
kind = "strip"
t_start = "0:10"
t_end = "0:20"
count = 3
caption = "Secuencia de tres fotogramas."
'''
        lab = load_lab(write_lab(self.dir, content))
        self.assertEqual([f.kind for f in lab.figures], ["pair", "strip"])
        self.assertEqual(len(lab.figures[0].frames), 2)

    def test_scaffold_creates_a_loadable_project(self):
        project = self.dir / "nuevo"
        scaffold(str(project))
        self.assertTrue((project / "content/report.toml").is_file())
        with self.assertRaises(FileExistsError):
            scaffold(str(project))


class StyleTests(unittest.TestCase):
    def messages(self, text, **kw):
        return " | ".join(i.message for i in lint_text(text, "x", **kw))

    def test_the_agent_style_that_graders_penalise_is_rejected(self):
        self.assertIn("etiqueta", self.messages("Advertencia Técnica Clave (Ámbito de Clases): el error fue anidar el método."))
        self.assertIn("«clave»", self.messages("Este ajuste es clave para el movimiento."))
        self.assertIn("sinergia", self.messages("Se logra una sinergia entre el audio y la física."))
        self.assertIn("muletilla", self.messages("En esta etapa se configura el componente."))
        self.assertIn("apreciaciones", self.messages("Creemos que el resultado es bueno."))

    def test_voseo_and_unfinished_sentences(self):
        self.assertIn("voseo", self.messages("Configurá el Rigidbody y ponelo en Interpolate."))
        self.assertNotIn("voseo", self.messages("Configura el Rigidbody."))
        self.assertIn("incompleta", self.messages("Configuramos el Collider para que este justo con la forma de"))

    def test_known_technical_mistakes(self):
        self.assertIn("fotograma", self.messages("El Update() se ejecuta cada milisegundo."))
        self.assertIn("GetKeyDown", self.messages("GetKey devuelve true una sola vez al presionar."))

    def test_plain_direct_text_passes(self):
        self.assertEqual(lint_text("Agregamos un Rigidbody al GameObject raíz de Mango.", "x"), [])
        self.assertEqual(lint_text("Componente Rigidbody en el Inspector", "x", sentence=False), [])

    def test_project_can_allow_a_word(self):
        self.assertEqual(lint_text("La clave foránea une las tablas.", "x", allow=["clave"]), [])

    def test_citations_must_exist_and_be_used(self):
        issues = check_citations({"a": "Texto [1] y [4]."}, 3)
        text = " ".join(i.message for i in issues)
        self.assertIn("[4] sin referencia", text)
        self.assertIn("[2] no se cita", text)
        self.assertIn("[3] no se cita", text)
        self.assertEqual(check_citations({"a": "[1] [2]"}, 2), [])

    @unittest.skipIf(Speller().dictionary is None, "requiere el diccionario hunspell es_ES")
    def test_spelling_catches_the_typos_of_the_hand_made_report(self):
        sp = Speller(["collider", "rigidbody", "input"])
        found = {i.message.split("«")[1].rstrip("»") for i in sp.check(
            "Agregmos un Colider. Ahora comprobamos que este reaccionando al scrip. El Input Hendler. Seleccionamos el componente.", "x")}
        self.assertEqual(found, {"Agregmos", "Colider", "scrip", "Hendler"})
        self.assertEqual(Speller(["assets", "audio"]).check("Creamos los archivos de la carpeta Assets/Audio y ajustamos el volumen de los pasos.", "x"), [])
        self.assertEqual(sp.check("Usamos `PlayOneShot` y GameObject con https://unity.com/xx.", "x"), [])


class LayoutTests(unittest.TestCase):
    def test_a_full_frame_is_rejected_when_its_text_would_be_unreadable(self):
        fit = fit_capture(1920, 1080, ui_text_px=15)
        self.assertIsNotNone(fit.problem)
        self.assertIn("recórtala", fit.problem)

    def test_scene_frames_are_not_checked_for_text(self):
        self.assertIsNone(fit_capture(1238, 594, ui_text_px=None, legible=False).problem)

    def test_tall_panels_never_exceed_the_height_cap(self):
        fit = fit_capture(437, 1000, ui_text_px=12)
        self.assertLessEqual(fit.height_in, MAX_H_IN + 1e-6)
        self.assertLessEqual(fit_capture(2000, 700, ui_text_px=None, legible=False).width_in, MAX_W_IN)

    def test_a_well_cropped_inspector_component_fits_with_readable_text(self):
        fit = fit_capture(437, 330, ui_text_px=12)
        self.assertIsNone(fit.problem)
        self.assertGreaterEqual(fit.ui_pt, 8)

    def test_code_images_keep_about_nine_points(self):
        fit = fit_code(700, 560)
        self.assertGreaterEqual(fit.ui_pt, 8)
        self.assertLessEqual(fit.height_in, 5.0 + 1e-6)
        self.assertIsNotNone(fit_code(1000, 1200).problem)          # demasiado alto: pide dividir

    def test_bitmap_is_downscaled_for_the_document(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "big.png"
            Image.new("RGB", (3000, 1500), "white").save(src)
            out = prepare_bitmap(src, Path(d) / "out" / "big", width_in=5.0)
            with Image.open(out) as im:
                self.assertEqual(im.width, 1100)                    # 5 in * 220 dpi
            jpg = prepare_bitmap(src, Path(d) / "out" / "game", width_in=5.0, photographic=True)
            self.assertEqual(jpg.suffix, ".jpg")


class SplitTests(unittest.TestCase):
    def test_parts_respect_the_cap_and_are_balanced(self):
        lines = CS.splitlines()
        parts = split_range(lines, 1, len(lines), 20)
        self.assertTrue(all(e - s + 1 <= 20 for s, e in parts))
        self.assertEqual(parts[0][0], 1)
        self.assertEqual(parts[-1][1], len(lines))
        sizes = [e - s + 1 for s, e in parts]
        self.assertGreaterEqual(min(sizes), max(sizes) // 2)          # sin trozos diminutos al final

    def test_no_tiny_trailing_part_for_22_lines(self):
        lines = [f"x{i}" for i in range(30)]
        parts = split_range(lines, 5, 26, 20)                        # 22 líneas
        self.assertEqual(len(parts), 2)
        self.assertGreaterEqual(min(e - s + 1 for s, e in parts), 8)

    def test_cuts_prefer_blank_lines_and_closing_braces(self):
        lines = ["a"] * 8 + ["}"] + ["b"] * 8 + [""] + ["c"] * 8
        parts = split_range(lines, 1, len(lines), 14)
        for s, e in parts[:-1]:
            self.assertTrue(lines[e - 1] in ("}", "") or lines[e] == "")

    def test_ranges_that_fit_stay_in_one_image_with_a_fold(self):
        groups = split_ranges(CS.splitlines(), [(1, 5), (10, 14)], 20)
        self.assertEqual(groups, [[(1, 5), (10, 14)]])

    def test_plan_code_numbers_images_with_real_lines(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            (root / "code").mkdir()
            (root / "code/Cat.cs").write_text(CS, encoding="utf-8")
            snips = plan_code(CodeSpec(id="todo", file="code/Cat.cs", max_lines=20, focus=[30]), root)
            self.assertGreater(len(snips), 1)
            self.assertEqual([s.id for s in snips][:2], ["todo-1", "todo-2"])
            self.assertEqual(sum(1 for s in snips if s.focus), 1)    # el resaltado va solo en el trozo que lo contiene
            with self.assertRaisesRegex(SnippetError, "enfocadas"):
                plan_code(CodeSpec(id="x", file="code/Cat.cs", lines="1-5", focus=[40]), root)


class CheckTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def test_minimal_lab_has_no_errors(self):
        rep = check_lab(load_lab(write_lab(self.dir, MINIMAL, video=False)), spell=False)
        self.assertEqual([str(i) for i in rep.errors], [])
        self.assertEqual(rep.metrics["imágenes de código"], 1)

    def test_bad_text_and_dangling_references_are_errors(self):
        content = MINIMAL.replace("Declaramos los campos de la clase.", "Es un paso clave. Mira la {fig:nada}.").replace(" [1]", "")
        rep = check_lab(load_lab(write_lab(self.dir, content, video=False)), spell=False)
        text = " ".join(str(i) for i in rep.errors)
        self.assertIn("«clave»", text)
        self.assertIn("{fig:nada}", text)
        self.assertIn("[1] no se cita", text)

    def test_visual_balance_warnings(self):
        content = MINIMAL.replace('[block.code]\nid = "campos"\nfile = "code/Cat.cs"\nlines = "5-14"\n', "")
        rep = check_lab(load_lab(write_lab(self.dir, content, video=False)), spell=False)
        self.assertTrue(any("sin figura ni código" in i.message for i in rep.warnings))
        self.assertTrue(any("ninguna imagen" in i.message for i in rep.errors))


@unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("tesseract") and has_chrome(), "requiere ffmpeg, tesseract y Chrome")
class EndToEndTests(unittest.TestCase):
    """Video sintético con texto grande -> figuras (OCR + recorte + anotación) + código -> .docx sobre la plantilla."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name) / "lab"
        cls.dir.mkdir()
        font = next((p for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                                 "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf") if Path(p).is_file()), None)
        if font is None:
            raise unittest.SkipTest("sin fuente para el video sintético")
        draw = ",".join([f"drawtext=fontfile={font}:text='Inspector':x=800:y=40:fontsize=56:fontcolor=white",
                         f"drawtext=fontfile={font}:text='Rigidbody':x=800:y=200:fontsize=56:fontcolor=white",
                         f"drawtext=fontfile={font}:text='Freeze Rotation':x=800:y=360:fontsize=56:fontcolor=white"])
        cls.video = cls.dir / "v.mp4"
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=0x383838:s=1280x720:d=6:r=10",
                        "-vf", draw, "-pix_fmt", "yuv420p", str(cls.video)], check=True)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_full_flow_builds_a_docx_with_captured_figures_and_code_images(self):
        content = MINIMAL + '''
[[block]]
kind = "step"
text = "Marcamos Freeze Rotation en el Rigidbody."
[block.figure]
id = "rigid"
t = 3
expect = ["Rigidbody", "Freeze Rotation"]
caption = "Componente Rigidbody."
[block.figure.crop]
around = "Rigidbody"
margin = [300, 250]
[block.figure.annotate]
target = "Freeze Rotation"
'''
        write_lab(self.dir, content)
        from src.unity_lab import build
        out = self.dir / "out"
        path = build(self.dir, str(out))
        document = docx.Document(str(path))
        table = document.tables[0]
        xml = table._tbl.xml
        self.assertEqual(xml.count("<w:drawing>"), 2)                    # figura del video + imagen de código
        texts = "\n".join(c.text for r in table.rows for c in r.cells)
        for expected in ("Práctica de prueba", "09", "01/01/2026", "Ana Pérez", "Luis Gómez", "Declaramos los campos",
                         "Figura 1:", "Figura 2:", "¿Qué hace Start?", "[1]"):
            self.assertIn(expected, texts)
        self.assertNotIn("Nombre Viejo", texts)
        self.assertNotIn("Ejercicio 1 del enunciado", texts)
        cant = [r for r in table.rows if r._tr.xpath("./w:trPr/w:cantSplit[@w:val='1']")]
        self.assertGreaterEqual(len(cant), 3)                            # un bloque por fila indivisible
        # figura anotada: contorno sin relleno y manifiesto verificable
        from src.evidence_engine import verify_manifest
        self.assertEqual([i for i in verify_manifest(str(out / "evidence")) if i.level == "error"], [])
        fig = Image.open(out / "figures" / "rigid.png")
        self.assertLessEqual(fig.width * fig.height, 1280 * 720)

    def test_wrong_expected_text_fails_with_a_review_sheet(self):
        content = MINIMAL + '''
[[block]]
kind = "step"
text = "Probamos el componente."
[block.figure]
id = "malo"
t = 3
expect = ["Spatial Blend"]
caption = "No aparece."
'''
        write_lab(self.dir, content)
        from src.unity_lab import FigureError, build
        with self.assertRaises(FigureError) as cm:
            build(self.dir, str(self.dir / "out2"))
        self.assertIn("Spatial Blend", str(cm.exception))
        self.assertTrue((self.dir / "out2" / "figures" / "_review" / "malo.png").is_file())

    def test_figures_are_cached_between_builds(self):
        content = MINIMAL + '''
[[block]]
kind = "step"
text = "Marcamos Freeze Rotation en el Rigidbody."
[block.figure]
id = "rigid"
t = 3
expect = ["Rigidbody"]
caption = "Componente Rigidbody."
[block.figure.crop]
around = "Rigidbody"
margin = [300, 250]
'''
        write_lab(self.dir, content)
        from src.unity_lab import build
        build(self.dir, str(self.dir / "out4"))
        with mock.patch("src.unity_lab.figures._capture") as capture:
            build(self.dir, str(self.dir / "out4"))
        capture.assert_not_called()                                      # misma especificación y mismo video: se reutiliza


class BlockRowsTests(unittest.TestCase):
    def test_each_block_gets_its_own_unsplittable_row_without_inner_borders(self):
        from src.docx_engine.block_rows import BlockRows
        with tempfile.TemporaryDirectory() as d:
            document = docx.Document(str(make_template(Path(d) / "t.docx")))
            table = document.tables[0]
            row = next(r for r in table.rows if r.cells[0].text.startswith("EJERCICIOS"))
            n = len(table.rows)

            def clear(cell):
                for p in cell.paragraphs[1:]:
                    p._p.getparent().remove(p._p)
                cell.paragraphs[0].text = ""
            rows = BlockRows(table, row, clear)
            for k in range(3):
                rows.next_cell().paragraphs[0].text = f"bloque {k}"
            rows.finish()
            self.assertEqual(len(table.rows), n + 2)
            blocks = [r for r in table.rows if r.cells[0].text.startswith("bloque")]
            self.assertEqual(len(blocks), 3)
            for r in blocks:
                self.assertTrue(r._tr.xpath("./w:trPr/w:cantSplit[@w:val='1']"))
            self.assertTrue(blocks[0]._tr.xpath(".//w:tcBorders/w:bottom[@w:val='nil']"))
            self.assertFalse(blocks[-1]._tr.xpath(".//w:tcBorders/w:bottom[@w:val='nil']"))
            self.assertTrue(blocks[1]._tr.xpath(".//w:tcBorders/w:top[@w:val='nil']"))


class DocxRenderTests(unittest.TestCase):
    def test_page_fill_detects_an_almost_empty_page(self):
        from src.core.docx_render import page_fill
        with tempfile.TemporaryDirectory() as d:
            empty, full = Path(d) / "e.png", Path(d) / "f.png"
            img = Image.new("RGB", (400, 560), "white")
            img.save(empty)
            from PIL import ImageDraw
            dr = ImageDraw.Draw(img)
            for y in range(110, 520, 5):
                dr.line([60, y, 340, y], fill="black", width=3)
            img.save(full)
            self.assertLess(page_fill(empty), 0.05)
            self.assertGreater(page_fill(full), 0.4)

    def test_clear_error_when_no_converter_is_available(self):
        from src.core import docx_render
        with mock.patch.object(docx_render.shutil, "which", return_value=None), \
                mock.patch.object(docx_render, "_flatpak_installed", return_value=False), \
                mock.patch.object(docx_render, "X2T_NATIVE", ()):
            with self.assertRaisesRegex(docx_render.RenderError, "libreoffice|LibreOffice"):
                docx_render.converter()


if __name__ == "__main__":
    unittest.main()
