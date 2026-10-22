# -*- coding: utf-8 -*-
"""Estilo sobrio de los .docx (igual que los informes Typst): sin azules, sin recuadros de color, pies de figura en negrita."""

import re
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image

from config import ReportConfig
from src.docx_engine import cell_helpers as ch


def _cell():
    doc = Document()
    return doc, doc.add_table(rows=1, cols=1).cell(0, 0)


def _runs(cell):
    return [r for p in cell.paragraphs for r in p.runs]


class DocxStyleTests(unittest.TestCase):
    def setUp(self):
        self.cfg = ReportConfig()

    def test_headings_and_bullets_use_the_text_color_not_institutional_blues(self):
        _, cell = _cell()
        ch.add_cell_heading_1(cell, "Etapa 1", self.cfg)
        ch.add_cell_heading_2(cell, "Subtítulo", self.cfg)
        ch.add_cell_bullet_point(cell, "Paso 1: ", "texto", self.cfg, level=1)
        ch.add_cell_bullet_point(cell, "Paso 1.1: ", "texto", self.cfg, level=2)
        colors = {str(r.font.color.rgb) for r in _runs(cell) if r.font.color and r.font.color.rgb is not None}
        self.assertEqual(colors, {str(ch.COLOR_BODY_TEXT)})
        symbols = [p.runs[0].text for p in cell.paragraphs if p.runs and p.runs[0].text.strip() in ("•", "–")]
        self.assertEqual([s.strip() for s in symbols], ["•", "–"])      # nada de ■ / ○ coloreados

    def test_code_block_is_light_gray_with_a_thin_border_all_around(self):
        _, cell = _cell()
        ch.add_cell_code_block(cell, "int x = 1;", self.cfg)
        xml = cell.paragraphs[-1]._p.xml
        self.assertIn('w:fill="F7F7F7"', xml)
        self.assertNotIn("1B365D", xml)
        for side in ("left", "top", "right", "bottom"):
            self.assertRegex(xml, rf'<w:{side} w:val="single" w:sz="4"')

    def test_figure_caption_has_bold_label_and_no_italics_or_grey(self):
        with tempfile.TemporaryDirectory() as d:
            img = Path(d) / "f.png"
            Image.new("RGB", (40, 20), "white").save(img)
            _, cell = _cell()
            ch.add_cell_image_with_caption(cell, str(img), "Figura 5: Código en el editor.", self.cfg)
        caption = [r for r in cell.paragraphs[-1].runs]
        self.assertEqual([(r.text, r.bold) for r in caption], [("Figura 5:", True), (" Código en el editor.", False)])
        self.assertTrue(all(not r.italic for r in caption))
        self.assertTrue(all(r.font.color.rgb == ch.COLOR_BODY_TEXT for r in caption))

    def test_callout_stays_neutral(self):
        from src.docx_engine.helpers import add_callout
        doc = Document()
        add_callout(doc, "Texto de la nota", title="Nota")
        xml = doc.tables[0]._tbl.xml
        colors = set(re.findall(r'w:(?:fill|color)="([0-9A-Fa-f]{6})"', xml))
        for c in colors:
            r, g, b = (int(c[i:i + 2], 16) for i in (0, 2, 4))
            self.assertLess(max(r, g, b) - min(r, g, b), 12, f"color con tinte: {c}")   # solo grises


if __name__ == "__main__":
    unittest.main()
