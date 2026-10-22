# -*- coding: utf-8 -*-
"""Imágenes de fragmentos de código: extracción, resaltado, spec, composición y captura real en Chrome."""

import os
import re
import tempfile
import unittest
from pathlib import Path

from PIL import Image

from src.core.chrome import BrowserError, find_chrome
from src.editor_engine import SnippetError, build_snippets, load_snippets
from src.editor_engine.extract import MAX_COLS, MAX_LINES, build_rows, parse_ranges, trim_blank_edges
from src.editor_engine.highlight import get_lexer, highlight_lines
from src.editor_engine.render import FONT_PX, LINE_PX, THEMES_DIR, size_of
from src.evidence_engine import manifest as mf
from src.evidence_engine.verify import verify_manifest
from src.slides_engine.validate import contrast_ratio

CS = """using UnityEngine;

public class Cat : MonoBehaviour
{
    public AudioClip jumpClip;   // clip de salto
    /* bloque
       multilínea */
    void Start()
    {
        var src = GetComponent<AudioSource>();
        src.volume = 0.5f;
        Debug.Log("listo");
    }
}
"""


def have_chrome() -> bool:
    try:
        find_chrome()
        return True
    except BrowserError:
        return False


def write_case(tmp: Path, body: str = "", code: str = CS) -> Path:
    (tmp / "Cat.cs").write_text(code, encoding="utf-8")
    spec = tmp / "snippets.toml"
    spec.write_text(body or '[[snippet]]\nid = "start"\nfile = "Cat.cs"\nlines = "8-13"\nfocus = [11]\n', encoding="utf-8")
    return spec


class ExtractTests(unittest.TestCase):
    def test_ranges(self):
        self.assertEqual(parse_ranges("12-30", 40), [(12, 30)])
        self.assertEqual(parse_ranges(7, 40), [(7, 7)])
        self.assertEqual(parse_ranges(["24-30", "12-18"], 40), [(12, 18), (24, 30)])
        self.assertEqual(parse_ranges(None, 5), [(1, 5)])

    def test_range_errors_are_reported_together(self):
        with self.assertRaises(SnippetError) as cm:
            parse_ranges(["0-3", "x", "30-99"], 40)
        self.assertEqual(len(cm.exception.problems), 3)
        with self.assertRaises(SnippetError):
            parse_ranges(["1-5", "4-8"], 10)

    def test_blank_edges_are_trimmed_but_numbering_is_kept(self):
        lines = ["", "a", "", "b", ""]
        self.assertEqual(trim_blank_edges(lines, [(1, 5)]), [(2, 4)])

    def test_limits(self):
        lines = [f"line {i}" for i in range(60)]
        with self.assertRaisesRegex(SnippetError, "máximo"):
            build_rows(lines, [(1, MAX_LINES + 1)], set(), False)
        with self.assertRaisesRegex(SnippetError, "columnas"):
            build_rows(["x" * (MAX_COLS + 1)], [(1, 1)], set(), False)
        with self.assertRaisesRegex(SnippetError, "enfocada"):
            build_rows(lines, [(1, 3)], {9}, False)

    def test_dedent_removes_common_indent_only(self):
        rows, cols = build_rows(["        a", "          b"], [(1, 2)], set(), True)
        self.assertEqual([r.text for r in rows], ["a", "  b"])
        self.assertEqual(cols, 3)

    def test_fold_row_between_ranges(self):
        rows, _ = build_rows([f"l{i}" for i in range(30)], [(1, 2), (10, 11)], set(), False)
        self.assertEqual([r.kind for r in rows], ["code", "code", "fold", "code", "code"])
        self.assertEqual([r.no for r in rows if r.kind == "code"], [1, 2, 10, 11])


class HighlightTests(unittest.TestCase):
    def classes(self):
        lines = highlight_lines(CS, get_lexer("Cat.cs", None))
        return {text.strip(): cls for line in lines for cls, text in line}, lines

    def test_csharp_tokens(self):
        cls, lines = self.classes()
        self.assertEqual(cls["using"], "kw")
        self.assertEqual(cls["Start"], "fn")
        self.assertEqual(cls["AudioClip"], "ty")
        self.assertEqual(cls["MonoBehaviour"], "ty")
        self.assertEqual(cls["0.5f"], "nu")
        self.assertEqual(cls['"listo"'], "st")
        self.assertEqual(cls["// clip de salto"], "cm")
        self.assertEqual(len(lines), len(CS.split("\n")))      # una lista de tokens por línea

    def test_multiline_comment_is_a_comment_on_each_line(self):
        _, lines = self.classes()
        self.assertTrue(all(c == "cm" for c, t in lines[5] if t.strip()))   # "/* bloque"
        self.assertTrue(all(c == "cm" for c, t in lines[6] if t.strip()))   # "multilínea */"

    def test_leading_blank_lines_do_not_shift_lines(self):
        lines = highlight_lines("\n\n\nx = 1\n", get_lexer("a.py", None))
        self.assertEqual("".join(t for _, t in lines[3]), "x = 1")
        self.assertEqual(lines[0], [])

    def test_unknown_language_falls_back_to_plain_text(self):
        lines = highlight_lines("hola mundo", get_lexer("notas.zzz", None))
        self.assertEqual("".join(t for _, t in lines[0]), "hola mundo")


class SpecTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def test_valid_spec_with_defaults(self):
        spec = write_case(self.dir, '[defaults]\ntheme = "light"\n[[snippet]]\nid = "a"\nfile = "Cat.cs"\nlines = "1-3"\n')
        snip = load_snippets(spec)[0]
        self.assertEqual((snip.id, snip.theme, snip.file.name), ("a", "light", "Cat.cs"))

    def test_all_problems_are_reported_together(self):
        spec = write_case(self.dir, '[[snippet]]\nid = "a"\nfile = "Cat.cs"\nlinez = "1-2"\n'
                                    '[[snippet]]\nid = "a"\nfile = "Falta.cs"\ntheme = "azul"\n'
                                    '[[snippet]]\nid = "b"\nfile = "Cat.cs"\ntext = "x"\n')
        with self.assertRaises(SnippetError) as cm:
            load_snippets(spec)
        text = str(cm.exception)
        for fragment in ("linez", "duplicado", "Falta.cs", "azul", "no ambos"):
            self.assertIn(fragment, text)

    def test_inline_snippet_needs_a_displayed_name(self):
        spec = write_case(self.dir, '[[snippet]]\nid = "a"\ntext = "x = 1"\n')
        with self.assertRaisesRegex(SnippetError, "title"):
            load_snippets(spec)

    def test_empty_spec(self):
        spec = write_case(self.dir, "# vacío\n")
        with self.assertRaisesRegex(SnippetError, r"\[\[snippet\]\]"):
            load_snippets(spec)


class ThemeAndLayoutTests(unittest.TestCase):
    def test_tokens_are_readable_in_both_themes(self):
        for theme in ("dark", "light"):
            css = (THEMES_DIR / f"{theme}.css").read_text(encoding="utf-8")
            v = dict(re.findall(r"--([\w-]+):\s*(#[0-9a-fA-F]{6})", css))
            for key in ("fg", "kw", "ty", "fn", "st", "nu", "cm", "at", "op", "bi", "pp"):
                self.assertGreaterEqual(contrast_ratio(v[key], v["bg"]), 4.5, f"{theme}:{key} sobre el fondo")
                self.assertGreaterEqual(contrast_ratio(v[key], v["line-hi"]), 4.0, f"{theme}:{key} sobre línea resaltada")
            self.assertGreaterEqual(contrast_ratio(v["no"], v["bg"]), 3.0, f"{theme}: números de línea")

    def test_light_theme_is_light_and_dark_is_dark(self):
        from src.slides_engine.validate import relative_luminance
        light = re.search(r"--bg:\s*(#\w{6})", (THEMES_DIR / "light.css").read_text()).group(1)
        dark = re.search(r"--bg:\s*(#\w{6})", (THEMES_DIR / "dark.css").read_text()).group(1)
        self.assertGreater(relative_luminance(light), 0.8)
        self.assertLess(relative_luminance(dark), 0.1)

    def test_themes_have_no_animation(self):
        for css in THEMES_DIR.glob("*.css"):
            text = re.sub(r"/\*.*?\*/", "", css.read_text(encoding="utf-8"), flags=re.S)
            self.assertNotRegex(text, r"animation|transition|@keyframes")

    def test_size_is_exact(self):
        rows, cols = build_rows(["a" * 20] * 5, [(1, 5)], set(), False)
        w, h = size_of(rows, cols, True)
        self.assertEqual(h, 42 + 34 + 14 + 5 * LINE_PX + 16)
        self.assertEqual(size_of(rows, cols, False)[1], 14 + 5 * LINE_PX + 16)
        self.assertGreaterEqual(w, 560)


def frame_width(first_line: int) -> int:
    rows, cols = build_rows(["x" * 60] * 3, [(1, 3)], set(), False)
    for k, r in enumerate(rows):
        r.no = first_line + k
    return size_of(rows, cols, True)[0]


class CompositionTests(unittest.TestCase):
    def frame_html(self, **kw):
        from src.editor_engine.api import _frame_and_meta
        from src.editor_engine.render import compose_html
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        spec = write_case(Path(tmp.name), kw.pop("body", ""))
        frame, meta = _frame_and_meta(load_snippets(spec)[0])
        return compose_html(frame), meta

    def test_html_is_static_and_uses_real_line_numbers(self):
        (html, w, h), meta = self.frame_html()
        self.assertNotRegex(html, r"<script|<link|https?://|animation|transition")
        self.assertIn("font-variant-ligatures: none", html)
        for n in range(8, 14):
            self.assertRegex(html, rf'class="no" style="width:[\d.]+px">{n}<')
        self.assertIn('class="ln hi"', html)
        self.assertEqual(meta["ranges"], [[8, 13]])
        self.assertEqual(meta["origin"], "file")

    def test_numbering_does_not_start_at_one(self):
        code = "\n".join(f"        int v{i} = {i};" for i in range(1, 140))
        spec = write_case(Path(self._tmp()), '[[snippet]]\nid = "x"\nfile = "Cat.cs"\nlines = "122-126"\nfocus = [124]\n', code=code)
        from src.editor_engine.api import _frame_and_meta
        from src.editor_engine.render import compose_html
        frame, meta = _frame_and_meta(load_snippets(spec)[0])
        html, _, _ = compose_html(frame)
        for n in range(122, 127):
            self.assertRegex(html, rf'class="no" style="width:[\d.]+px">{n}<')
        self.assertNotRegex(html, r'class="no" style="width:[\d.]+px">1<')
        self.assertEqual(meta["ranges"], [[122, 126]])
        text = re.sub(r"<[^>]+>", "", html)
        self.assertIn("int v122 = 122;", text)          # el código mostrado corresponde a esas líneas
        self.assertIn("int v126 = 126;", text)
        self.assertNotIn("int v121 =", text)

    def test_inline_snippet_can_declare_its_first_line_number(self):
        spec = write_case(Path(self._tmp()), '[[snippet]]\nid = "i"\ntitle = "a.py"\nstart = 1048\ntext = "x = 1\\ny = 2"\n')
        from src.editor_engine.api import _frame_and_meta
        from src.editor_engine.render import compose_html
        frame, meta = _frame_and_meta(load_snippets(spec)[0])
        html, _, _ = compose_html(frame)
        self.assertRegex(html, r'>1048<')
        self.assertRegex(html, r'>1049<')
        self.assertEqual(meta["ranges"], [[1048, 1049]])
        text = re.sub(r"<[^>]+>", "", html)
        self.assertIn("x = 1", text)
        self.assertIn("y = 2", text)
        # con 4 dígitos el margen se ensancha: la imagen debe ser más ancha que con 3
        self.assertGreater(frame_width(1048), frame_width(48))

    def test_start_is_rejected_with_a_real_file(self):
        spec = write_case(Path(self._tmp()), '[[snippet]]\nid = "x"\nfile = "Cat.cs"\nstart = 5\n')
        with self.assertRaisesRegex(SnippetError, "start"):
            load_snippets(spec)

    def _tmp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        return tmp.name

    def test_fold_marker_between_ranges(self):
        body = '[[snippet]]\nid = "x"\nfile = "Cat.cs"\nlines = ["1-1", "8-9"]\n'
        (html, _, _), meta = self.frame_html(body=body)
        self.assertIn("⋯", html)
        self.assertEqual(meta["ranges"], [[1, 1], [8, 9]])

    def test_illegible_width_is_rejected(self):
        wide = "a" * 85
        body = f'[[snippet]]\nid = "w"\ntitle = "w.py"\ntext = """{wide}"""\n'
        from src.editor_engine.api import _frame_and_meta
        with tempfile.TemporaryDirectory() as tmp:
            snip = load_snippets(write_case(Path(tmp), body))[0]
            with self.assertRaisesRegex(SnippetError, "pt"):
                _frame_and_meta(snip)


@unittest.skipUnless(have_chrome(), "requiere Chrome/Chromium")
class RealBrowserTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)
        self.ev = self.dir / "evidence"

    def test_png_has_exact_size_content_and_manifest(self):
        spec = write_case(self.dir)
        item = build_snippets(spec, out_dir=self.ev)[0]
        img = Image.open(self.ev / item["image"])
        self.assertEqual(img.size, (item["width_px"] * 2, item["height_px"] * 2))
        self.assertGreater(len(set(img.convert("RGB").resize((64, 64)).getdata())), 8)   # no está en blanco
        self.assertEqual(img.getpixel((img.width // 2, img.height - 2)), img.getpixel((5, img.height - 2)))   # sin recorte raro
        self.assertEqual(item["source"], "code")
        self.assertEqual(item["render"], "editor-frame")
        self.assertEqual([i for i in verify_manifest(str(self.ev)) if i.level == "error"], [])

    def test_light_theme_has_a_light_background(self):
        spec = write_case(self.dir, '[[snippet]]\nid = "l"\nfile = "Cat.cs"\nlines = "8-13"\ntheme = "light"\n')
        item = build_snippets(spec, out_dir=self.ev)[0]
        img = Image.open(self.ev / item["image"]).convert("RGB")
        self.assertGreater(sum(img.getpixel((img.width // 2, img.height - 6))), 600)

    def test_editing_the_source_is_detected(self):
        spec = write_case(self.dir)
        build_snippets(spec, out_dir=self.ev)
        (self.dir / "Cat.cs").write_text(CS.replace("0.5f", "0.9f"), encoding="utf-8")
        errors = [i for i in verify_manifest(str(self.ev)) if i.level == "error"]
        self.assertTrue(any("fuente cambió" in i.message for i in errors), errors)

    def test_editing_the_image_is_detected(self):
        spec = write_case(self.dir)
        item = build_snippets(spec, out_dir=self.ev)[0]
        with open(self.ev / item["image"], "ab") as fh:
            fh.write(b"x")
        self.assertTrue(any("modificado" in i.message for i in verify_manifest(str(self.ev)) if i.level == "error"))

    def test_inline_snippet_is_registered_without_warnings(self):
        spec = write_case(self.dir, '[[snippet]]\nid = "i"\ntitle = "demo.py"\ntext = "x = 1\\nprint(x)"\n')
        item = build_snippets(spec, out_dir=self.ev)[0]
        self.assertEqual(item["origin"], "inline")
        self.assertEqual(verify_manifest(str(self.ev)), [])      # es contenido del reporte: no hay fuente con que compararlo

    def test_docx_helper_sizes_text_consistently(self):
        from docx import Document
        from docx.shared import Inches
        from src.docx_engine.cell_helpers import add_cell_code_figure
        spec = write_case(self.dir)
        build_snippets(spec, out_dir=self.ev)
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)

        class Cfg:
            font_family_body, font_size_caption_pt = "Calibri", 9
        add_cell_code_figure(cell, "start", self.ev, "Figura 1: prueba", Cfg())
        width_in = doc.inline_shapes[0].width / 914400
        self.assertLessEqual(width_in, 6.0)
        with self.assertRaises(KeyError):
            add_cell_code_figure(cell, "no-existe", self.ev, "x", Cfg())

    def test_slides_use_the_image_through_evidence_id(self):
        from src.slides_engine import render_deck, load_deck
        spec = write_case(self.dir, '[[snippet]]\nid = "start"\nfile = "Cat.cs"\nlines = "8-13"\ntheme = "light"\n')
        build_snippets(spec, out_dir=self.ev)
        proj = self.dir / "deck"
        (proj / "content").mkdir(parents=True)
        (proj / "project.toml").write_text('pipeline = "html-slides"\n[slides]\ntitle = "Demo"\nauthors = ["A, B"]\n', encoding="utf-8")
        (proj / "content/slides.toml").write_text(
            '[[slide]]\nlayout = "title"\ntitle = "Demo"\n\n[[slide]]\nlayout = "figure"\ntitle = "Inicio"\n'
            'evidence_id = "start"\ncaption = "**Figura 1.** Fragmento."\n', encoding="utf-8")
        from src.reports.html_slides import build
        out = build(proj, self.dir / "out", evidence_dir=self.ev)
        self.assertTrue(Path(out).is_file())


if __name__ == "__main__":
    unittest.main()
