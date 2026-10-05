#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script Maestro de Generación del Informe de Negocios Electrónicos (Lab 05)
Distribuidora Inca S.R.L. - Solución SCM con Odoo Community Edition
Universidad Nacional de San Agustín de Arequipa (UNSA) - EPIS
Docente: Dr. Ing. César Basilio Baluarte Araya
"""

import os
import sys
import docx

from build_sections_part1 import (
    build_portada_and_indice,
    build_section_1_planificacion,
    build_section_2_organizacion,
    build_section_3_problema,
    build_section_4_marco_teorico
)

from build_sections_part2 import (
    build_section_5_comparativas,
    build_section_6_trabajo_equipo,
    build_section_7_alternativas,
    build_section_8_seleccion,
    build_section_9_presentacion,
    build_section_10_prototipo,
    build_section_11_lecciones,
    build_section_12_conclusiones,
    build_section_13_referencias,
    build_section_14_anexos,
    build_section_15_informe,
    build_section_16_autoevaluacion
)

def main():
    template_path = "inputs/NE/Lab05/GenericTemplate.docx"
    output_dir = "output/NE/Lab05"
    output_path = os.path.join(output_dir, "Informe_Lab05_Distribuidora_Inca_SCM.docx")

    print(f"[*] Cargando plantilla oficial: {template_path}")
    if not os.path.exists(template_path):
        print(f"[!] Error: No se encontró la plantilla en {template_path}")
        sys.exit(1)

    doc = docx.Document(template_path)
    os.makedirs(output_dir, exist_ok=True)

    print("[*] Construyendo Portada Institucional e Índice General...")
    build_portada_and_indice(doc)

    print("[*] Construyendo Sección 1: Planificar el tratamiento del problema...")
    build_section_1_planificacion(doc)

    print("[*] Construyendo Sección 2: Organizar el trabajo del equipo/grupo...")
    build_section_2_organizacion(doc)

    print("[*] Construyendo Sección 3: Problema (Contexto, Enunciado y Árbol de Causas)...")
    build_section_3_problema(doc)

    print("[*] Construyendo Sección 4: Marco Teórico (10 conceptos, 4 tesis, 4 artículos, arquitectura)...")
    build_section_4_marco_teorico(doc)

    print("[*] Construyendo Sección 5: Comparativa de Aspectos (Modelos, Métodos, Técnicas y Herramientas)...")
    build_section_5_comparativas(doc)

    print("[*] Construyendo Sección 6: Trabajo Colaborativo en Equipo (Matriz RACI)...")
    build_section_6_trabajo_equipo(doc)

    print("[*] Construyendo Sección 7: Generación de Posibles Soluciones (Alternativas 1, 2 y 3)...")
    build_section_7_alternativas(doc)

    print("[*] Construyendo Sección 8: Selección de Alternativa y Presupuesto Detallado (S/. 6,000.00)...")
    build_section_8_seleccion(doc)

    print("[*] Construyendo Sección 9: Presentación de la Solución (PPT, Video, Nomenclatura)...")
    build_section_9_presentacion(doc)

    print("[*] Construyendo Sección 10: Prototipo Odoo (4 Flujos SCM y 12 Placeholders de Captura)...")
    build_section_10_prototipo(doc)

    print("[*] Construyendo Sección 11: Lecciones Aprendidas (Estructura Dual)...")
    build_section_11_lecciones(doc)

    print("[*] Construyendo Sección 12: Conclusiones (Alineadas 1 a 1 con Objetivos)...")
    build_section_12_conclusiones(doc)

    print("[*] Construyendo Sección 13: Referencias Bibliográficas (Norma IEEE)...")
    build_section_13_referencias(doc)

    print("[*] Construyendo Sección 14: Anexos (Láminas PPT, Infografía SCM y Dataset Maestro)...")
    build_section_14_anexos(doc)

    print("[*] Construyendo Sección 15: Informe (Expresión Escrita y Coherencia)...")
    build_section_15_informe(doc)

    print("[*] Construyendo Sección 16: Autoevaluación Individual del Equipo...")
    build_section_16_autoevaluacion(doc)

    print(f"[*] Guardando documento final en: {output_path}")
    doc.save(output_path)
    
    file_size = os.path.getsize(output_path)
    print(f"[+] ¡INFORME GENERADO CON ÉXITO! Tamaño del archivo: {file_size / (1024*1024):.2f} MB ({file_size:,} bytes)")
    print(f"[+] Total de párrafos en el documento: {len(doc.paragraphs)}")
    print(f"[+] Total de tablas en el documento: {len(doc.tables)}")

if __name__ == "__main__":
    main()
